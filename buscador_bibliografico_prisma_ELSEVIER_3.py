#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Buscador bibliográfico + diagrama de flujo PRISMA 2020  (interfaz gráfica)

Desarrolladores: Paul Prudencio Galvez y Marcelino Prudencio Chanca

Sube un Word (.docx) con su manuscrito. El programa:
  1. Lee título, resumen, palabras clave y los DOI que ya cita.
  2. Propone consultas de búsqueda (usted puede editarlas).
  3. Busca en: Google Scholar*, SciELO**, DOAJ, Crossref, SpringerOpen***, PubMed,
     ERIC, Redalyc**, RENATI**, La Referencia**, OpenAlex(a), BASE(b), Dimensions(c),
     JISC Library Hub Discover**, KERWA (Univ. de Costa Rica)**, ELSEVIER(d), arXiv, Semantic Scholar(e),
     CORE(f), Europe PMC, OpenCitations(g), Zenodo, Figshare, Dryad, DataCite, Lens.org(h), OpenAIRE, DOAB,
     OAPEN, PLOS, eLife y, solo como enlace manual, Scilit y WorldCat (no tienen API pública).
     DOI.org se usa para VALIDAR cada DOI.  Idioma: puede buscar en INGLÉS, ESPAÑOL o ambos (botón para
     añadir las consultas traducidas al inglés y filtro por idioma del documento).
  4. Elimina duplicados, filtra por años, descarta lo ya citado, puntúa la relevancia
     y valida los DOI (doi.org + Crossref).
  5. Genera UN SOLO ARCHIVO DE WORD EDITABLE con todos los resultados: consultas, registros por
     repositorio, diagrama de flujo PRISMA 2020 EDITABLE en ESPAÑOL e INGLÉS (estructura Identificación /
     Cribado: base de datos, cadena de búsqueda, filtros, recolección de datos, duplicados, cribado y
     elegibilidad), cuadro por hallazgos (grupos temáticos y tipo de evidencia), referencias detalladas
     y cribado completo.
  6. Botones: «Nueva consulta» (limpia la ventana para otra búsqueda) y «Salir».

  (a) OpenAlex exige una clave GRATUITA desde el 13/02/2026 (openalex.org/settings/api).
  (b) BASE solo responde a IP autorizadas por BASE (solicítelo en base-search.net).
  (c) Dimensions requiere credenciales de API (suscripción o acceso institucional).
  (e) Semantic Scholar funciona sin clave, pero con límite compartido; una clave gratuita lo estabiliza.
  (f) CORE exige una clave gratuita (core.ac.uk/services/api).
  (g) OpenCitations es un índice de CITAS, no un buscador temático: usa como semilla los DOI que su manuscrito
      ya cita y recupera las obras que los citan (se completa con Crossref).
  (h) Lens.org exige un token de API (solicitud en lens.org).
  (d) Elsevier: con clave gratuita de dev.elsevier.com se usa la API de ScienceDirect Search v2; sin clave se
      consulta Crossref filtrando el prefijo DOI 10.1016 (Elsevier), marcado como tal en el reporte.

  *  Google Scholar no tiene API oficial: se consulta solo si usted ingresa una clave de
     SerpAPI (serpapi.com). Sin clave se omite y se deja el enlace de búsqueda manual.
  ** SciELO, Redalyc, RENATI y La Referencia no ofrecen una API de búsqueda estable:
     se consultan "a mejor esfuerzo". Si una falla, el reporte lo registra y deja el
     enlace para buscar a mano. No se inventan resultados.
  *** SpringerOpen usa la API de Springer Nature si ingresa su clave (gratuita); sin clave
     se consulta Crossref filtrando el prefijo DOI 10.1186 (BMC/SpringerOpen).

Requisitos:  pip install python-docx
             (tkinter viene con Python; en Linux: sudo apt install python3-tk)
Uso:         python buscador_bibliografico_prisma_ELSEVIER.py

Aviso metodológico: el cribado es AUTOMÁTICO (relevancia léxica y validación de DOI). No
sustituye el cribado por dos revisores que exige una revisión sistemática; el diagrama
sigue la estructura PRISMA 2020 para búsquedas en bases de datos.
"""

import re
import os
import html
import tempfile
import json
import time
import queue
import difflib
import threading
import urllib.request
import urllib.error
import urllib.parse
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

# ============================================================== constantes
UA = "buscador-bibliografico-prisma/1.0"
DOI_RE = re.compile(r"10\.\d{4,9}/[^\s\"<>]+", re.IGNORECASE)
ANIO_ACTUAL = datetime.now().year
DESARROLLADORES = "Paul Prudencio Galvez y Marcelino Prudencio Chanca"

STOP = set("""a al algo ante con contra cual cuando como de del desde donde el ella ellas ellos en entre era es esta
estas este estos la las lo los más mas muy no nos o para pero por que se sin sobre su sus también tambien un una uno
unas unos y ya the and for with from that this these those are was were been into using use used based between among
within over under their than then which what when where who how can may more most such also not but its has have had
each other some any all one two three our your his her they them there here about after before during hacia segun
según través mediante partir caso casos estudio estudios análisis analisis revisión revision artículo articulo
modelo modelos nuevo nueva nuevos nuevas través""".split())


# ============================================================== utilidades
def norm(s):
    return re.sub(r"[^a-z0-9áéíóúñü ]+", " ", (s or "").lower()).strip()


def tokens(s, minlen=4):
    s = re.sub(r"<[^>]+>", " ", (s or "").lower())
    return {t for t in re.findall(r"[a-záéíóúñü0-9]+", s) if len(t) >= minlen and t not in STOP}


def limpiar_doi(d):
    d = (d or "").strip().rstrip(".,;:)]}>\"'&")
    d = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", d, flags=re.I)
    if d.count("(") > d.count(")"):
        d += ")"
    return d.lower()


def http_get(url, headers=None, timeout=30, reintentos=2):
    h = {"User-Agent": UA, "Accept": "application/json, text/html;q=0.8, */*;q=0.5"}
    if headers:
        h.update(headers)
    ultimo = None
    for i in range(reintentos + 1):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=h), timeout=timeout) as r:
                return r.status, r.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as e:
            if e.code in (400, 401, 403, 404, 409):
                return e.code, ""
            ultimo = e
        except Exception as e:
            ultimo = e
        time.sleep(1.2 * (i + 1))
    raise RuntimeError(f"sin respuesta ({ultimo})")


def http_post(url, data, headers=None, timeout=40):
    h = {"User-Agent": UA}
    if headers:
        h.update(headers)
    body = data.encode("utf-8") if isinstance(data, str) else data
    try:
        with urllib.request.urlopen(urllib.request.Request(url, data=body, headers=h, method="POST"), timeout=timeout) as r:
            return r.status, r.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception as e:
        raise RuntimeError(f"sin respuesta ({e})")


def get_json(url, headers=None):
    st, txt = http_get(url, headers)
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    return json.loads(txt)


def registro(**kw):
    base = dict(titulo="", autores="", anio=None, revista="", doi="", url="", resumen="", fuente="", consulta="",
                meta_crossref=False)
    base.update(kw)
    base["doi"] = limpiar_doi(base["doi"]) if base["doi"] else ""
    return base


def anio_de(texto):
    m = re.search(r"\b(19|20)\d{2}\b", str(texto or ""))
    return int(m.group(0)) if m else None


# ============================================================== lectura del manuscrito
def leer_parrafos(ruta):
    doc = Document(ruta)
    out = [p.text.strip() for p in doc.paragraphs]
    for t in doc.tables:
        for f in t.rows:
            for c in f.cells:
                out += [p.text.strip() for p in c.paragraphs]
    return [t for t in out if t]


def perfil_manuscrito(parrafos):
    titulo = ""
    for p in parrafos[:15]:
        if 25 <= len(p) <= 300 and not re.match(r"(manuscrito|highlights|resumen|abstract)", p, re.I):
            titulo = p
            break
    kws = []
    for p in parrafos:
        m = re.match(r"^\s*(palabras\s+clave|keywords?|key\s+words)\s*[:：]\s*(.+)$", p, re.I)
        if m:
            kws += [k.strip(" .") for k in re.split(r"[;,]", m.group(2)) if k.strip(" .")]
    resumen = ""
    for i, p in enumerate(parrafos):
        if re.fullmatch(r"(resumen|abstract)", p.strip(), re.I) and i + 1 < len(parrafos):
            resumen += " " + parrafos[i + 1]
    dois = {limpiar_doi(d) for p in parrafos for d in DOI_RE.findall(p)}
    base = " ".join([titulo, resumen] + kws)
    freq = Counter(t for t in re.findall(r"[a-záéíóúñü]{5,}", base.lower()) if t not in STOP)
    terminos = [t for t, _ in freq.most_common(12)]
    return dict(titulo=titulo, keywords=kws, resumen=resumen.strip(), dois_citados=dois, terminos=terminos)


def construir_consultas(perfil, maximo=10):
    qs = []
    if perfil["titulo"]:
        qs.append(" ".join(perfil["titulo"].split()[:14]))
    k = perfil["keywords"]
    qs += [f"{k[i]} {k[i + 1]}" for i in range(0, len(k) - 1, 2)]
    qs += k[:4]
    if not k:
        t = perfil["terminos"]
        qs += [" ".join(t[:4]), " ".join(t[4:8])]
    vistos, out = set(), []
    for q in qs:
        n = norm(q)
        if n and n not in vistos:
            vistos.add(n)
            out.append(q.strip())
    return out[:maximo]


# ============================================================== conectores
def _cr_registro(it, fuente):
    f = lambda a: (", ".join(x for x in (a.get("family", ""), a.get("given", "")) if x)) or a.get("name", "")
    anio = None
    for k in ("issued", "published-print", "published-online"):
        dp = (it.get(k) or {}).get("date-parts")
        if dp and dp[0] and dp[0][0]:
            anio = int(dp[0][0])
            break
    return registro(titulo=((it.get("title") or [""])[0]).strip(),
                    autores="; ".join(filter(None, (f(a) for a in it.get("author", [])))),
                    anio=anio, revista=(it.get("container-title") or [""])[0], doi=it.get("DOI", ""),
                    url=it.get("URL", ""), resumen=re.sub(r"<[^>]+>", " ", it.get("abstract", "") or ""),
                    fuente=fuente, meta_crossref=True)


def buscar_crossref(q, n, y0, y1, cfg, prefijo=None, fuente="Crossref"):
    filtro = f"from-pub-date:{y0},until-pub-date:{y1},type:journal-article"
    if prefijo:
        filtro += f",prefix:{prefijo}"
    par = {"query.bibliographic": q, "rows": n, "filter": filtro,
           "select": "DOI,title,author,issued,container-title,abstract,URL"}
    if cfg.get("correo"):
        par["mailto"] = cfg["correo"]
    d = get_json("https://api.crossref.org/works?" + urllib.parse.urlencode(par))
    return [_cr_registro(it, fuente) for it in d["message"]["items"]]


def buscar_pubmed(q, n, y0, y1, cfg):
    base = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/"
    par = {"db": "pubmed", "term": q + {"en": " AND eng[lang]", "es": " AND spa[lang]"}.get(cfg.get("idioma"), ""), "retmax": n, "retmode": "json", "datetype": "pdat",
           "mindate": y0, "maxdate": y1}
    ids = get_json(base + "esearch.fcgi?" + urllib.parse.urlencode(par))["esearchresult"].get("idlist", [])
    if not ids:
        return []
    time.sleep(0.4)
    st, xml = http_get(base + "efetch.fcgi?" + urllib.parse.urlencode({"db": "pubmed", "id": ",".join(ids), "retmode": "xml"}))
    out = []
    for a in ET.fromstring(xml).findall(".//PubmedArticle"):
        t = a.find(".//ArticleTitle")
        doi = ""
        for i in a.findall(".//ArticleIdList/ArticleId"):
            if i.get("IdType") == "doi":
                doi = i.text or ""
        pm = a.findtext(".//MedlineCitation/PMID", "")
        out.append(registro(
            titulo="".join(t.itertext()).strip() if t is not None else "",
            autores="; ".join(f"{x.findtext('LastName', '')}, {x.findtext('ForeName', '')}".strip(", ")
                              for x in a.findall(".//AuthorList/Author") if x.findtext("LastName")),
            anio=anio_de(a.findtext(".//JournalIssue/PubDate/Year") or a.findtext(".//JournalIssue/PubDate/MedlineDate")),
            revista=a.findtext(".//Journal/Title", ""), doi=doi,
            url=f"https://pubmed.ncbi.nlm.nih.gov/{pm}/" if pm else "",
            resumen=" ".join("".join(x.itertext()) for x in a.findall(".//Abstract/AbstractText")), fuente="PubMed"))
    time.sleep(0.4)
    return out


def buscar_doaj(q, n, y0, y1, cfg):
    d = get_json(f"https://doaj.org/api/search/articles/{urllib.parse.quote(q, safe='')}?pageSize={n}")
    out = []
    for r in d.get("results", []):
        b = r.get("bibjson", {})
        doi = next((i["id"] for i in b.get("identifier", []) if i.get("type", "").lower() == "doi"), "")
        url = next((l["url"] for l in b.get("link", []) if l.get("url")), "")
        out.append(registro(titulo=b.get("title", ""), autores="; ".join(a.get("name", "") for a in b.get("author", [])),
                            anio=anio_de(b.get("year")), revista=(b.get("journal") or {}).get("title", ""),
                            doi=doi, url=url, resumen=b.get("abstract", ""), fuente="DOAJ"))
    return out


def buscar_eric(q, n, y0, y1, cfg):
    par = {"search": q, "format": "json", "rows": n, "start": 0}
    d = get_json("https://api.ies.ed.gov/eric/?" + urllib.parse.urlencode(par))
    out = []
    for r in d.get("response", {}).get("docs", []):
        doi = r.get("doi") or ""
        if not doi:
            m = DOI_RE.search(json.dumps(r))
            doi = m.group(0) if m else ""
        au = r.get("author", [])
        out.append(registro(titulo=r.get("title", ""), autores="; ".join(au) if isinstance(au, list) else str(au),
                            anio=anio_de(r.get("publicationdateyear")), revista=r.get("source", ""), doi=doi,
                            url=r.get("url") or f"https://eric.ed.gov/?id={r.get('id', '')}",
                            resumen=r.get("description", ""), fuente="ERIC"))
    return out


def buscar_springeropen(q, n, y0, y1, cfg):
    if cfg.get("key_springer"):
        par = {"q": q, "p": n, "api_key": cfg["key_springer"]}
        d = get_json("https://api.springernature.com/openaccess/json?" + urllib.parse.urlencode(par))
        out = []
        for r in d.get("records", []):
            url = (r.get("url") or [{}])[0].get("value", "") if isinstance(r.get("url"), list) else ""
            out.append(registro(titulo=r.get("title", ""),
                                autores="; ".join(c.get("creator", "") for c in r.get("creators", [])),
                                anio=anio_de(r.get("publicationDate")), revista=r.get("publicationName", ""),
                                doi=r.get("doi", ""), url=url, resumen=r.get("abstract", ""), fuente="SpringerOpen"))
        return out
    return buscar_crossref(q, n, y0, y1, cfg, prefijo="10.1186", fuente="SpringerOpen (Crossref, prefijo 10.1186)")


def _doi_por_titulo(titulo, cfg):
    """Busca el DOI de un título en Crossref; solo acepta coincidencias muy altas."""
    par = {"query.bibliographic": titulo, "rows": 1, "select": "DOI,title,author,issued,container-title,abstract,URL"}
    if cfg.get("correo"):
        par["mailto"] = cfg["correo"]
    try:
        it = get_json("https://api.crossref.org/works?" + urllib.parse.urlencode(par))["message"]["items"][0]
    except Exception:
        return None
    r = _cr_registro(it, "")
    return r if difflib.SequenceMatcher(None, norm(r["titulo"]), norm(titulo)).ratio() >= 0.85 else None


def buscar_scholar(q, n, y0, y1, cfg):
    if not cfg.get("key_serpapi"):
        raise RuntimeError("omitido: sin clave de SerpAPI (Google Scholar no tiene API oficial)")
    par = {"engine": "google_scholar", "q": q, "num": min(n, 20), "as_ylo": y0, "as_yhi": y1, "api_key": cfg["key_serpapi"]}
    d = get_json("https://serpapi.com/search.json?" + urllib.parse.urlencode(par))
    out = []
    for r in d.get("organic_results", []):
        resumen = (r.get("publication_info") or {}).get("summary", "")
        rec = registro(titulo=r.get("title", ""), autores=resumen.split(" - ")[0], anio=anio_de(resumen),
                       revista=resumen.split(" - ")[1].split(",")[0] if resumen.count(" - ") >= 1 else "",
                       url=r.get("link", ""), resumen=r.get("snippet", ""), fuente="Google Scholar")
        cr = _doi_por_titulo(rec["titulo"], cfg)
        if cr:
            rec.update(doi=cr["doi"], autores=cr["autores"] or rec["autores"], revista=cr["revista"] or rec["revista"],
                       anio=cr["anio"] or rec["anio"], meta_crossref=True)
        out.append(rec)
        time.sleep(0.2)
    return out


def _dois_en_html(url, fuente, n, cfg):
    """Portales sin API: se leen los DOI que aparecen en la página de resultados y se resuelven en Crossref."""
    st, html = http_get(url)
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    dois = []
    for d in DOI_RE.findall(html):
        d = limpiar_doi(re.split(r"[&\\'\"<]", d)[0])
        if d and d not in dois:
            dois.append(d)
    out = []
    for d in dois[:n]:
        try:
            it = get_json("https://api.crossref.org/works/" + urllib.parse.quote(d, safe="/()"))["message"]
            r = _cr_registro(it, fuente)
            out.append(r)
        except Exception:
            pass
        time.sleep(0.15)
    return out


def buscar_scielo(q, n, y0, y1, cfg):
    return _dois_en_html("https://search.scielo.org/?" + urllib.parse.urlencode({"lang": "es", "count": n, "q": q}),
                         "SciELO", n, cfg)


def buscar_redalyc(q, n, y0, y1, cfg):
    return _dois_en_html("https://www.redalyc.org/busquedaArticuloFiltros.oa?" + urllib.parse.urlencode({"q": q}),
                         "Redalyc", n, cfg)


def buscar_lareferencia(q, n, y0, y1, cfg):
    par = [("lookfor", q), ("type", "AllFields"), ("limit", n), ("field[]", "title"), ("field[]", "authors"),
           ("field[]", "publicationDates"), ("field[]", "urls"), ("field[]", "summary"), ("field[]", "id")]
    d = get_json("https://www.lareferencia.info/vufind/api/v1/search?" + urllib.parse.urlencode(par))
    out = []
    for r in d.get("records", []):
        au = r.get("authors", {})
        nombres = []
        if isinstance(au, dict):
            for k in ("primary", "secondary"):
                v = au.get(k) or {}
                nombres += list(v.keys()) if isinstance(v, dict) else [str(x) for x in v]
        elif isinstance(au, list):
            nombres = [str(x) for x in au]
        au = nombres
        urls = r.get("urls") or []
        url = urls[0].get("url", "") if urls and isinstance(urls[0], dict) else ""
        txt = json.dumps(r, ensure_ascii=False)
        m = DOI_RE.search(txt)
        out.append(registro(titulo=r.get("title", ""), autores="; ".join(au) if isinstance(au, list) else "",
                            anio=anio_de(r.get("publicationDates")), url=url, doi=m.group(0) if m else "",
                            resumen=" ".join(r.get("summary", [])) if isinstance(r.get("summary"), list) else "",
                            fuente="La Referencia"))
    return out


def _dspace7(base, fuente, etiqueta, q, n):
    """Repositorios DSpace 7 (RENATI, KERWA): API REST pública de descubrimiento."""
    d = get_json(base + "/server/api/discover/search/objects?" +
                 urllib.parse.urlencode({"query": q, "size": n, "dsoType": "ITEM"}))
    objs = d["_embedded"]["searchResult"]["_embedded"]["objects"]
    out = []
    for o in objs:
        it = o["_embedded"]["indexableObject"]
        md = it.get("metadata", {})
        g = lambda k: [x.get("value", "") for x in md.get(k, [])]
        uri = (g("dc.identifier.uri") or [""])[0]
        doi = next((x for x in g("dc.identifier.doi")), "")
        if not doi:
            m = DOI_RE.search(" ".join(g("dc.identifier.uri") + g("dc.relation.uri") + g("dc.identifier.other")))
            doi = m.group(0) if m else ""
        rev = (g("dc.relation.ispartof") or g("dc.source") or [etiqueta])[0]
        out.append(registro(titulo=(g("dc.title") or [it.get("name", "")])[0], autores="; ".join(g("dc.contributor.author")),
                            anio=anio_de((g("dc.date.issued") or [""])[0]), revista=rev, url=uri, doi=doi,
                            resumen=" ".join(g("dc.description.abstract")), fuente=fuente))
    return out


def buscar_renati(q, n, y0, y1, cfg):
    return _dspace7("https://renati.sunedu.gob.pe", "RENATI", "Tesis / trabajo académico (RENATI)", q, n)


def buscar_kerwa(q, n, y0, y1, cfg):
    """KÉRWÁ: repositorio institucional de la Universidad de Costa Rica (DSpace 7.6)."""
    return _dspace7("https://www.kerwa.ucr.ac.cr", "KERWA", "Repositorio Kérwá (Universidad de Costa Rica)", q, n)


def _abstract_openalex(inv):
    if not inv:
        return ""
    pos = {}
    for palabra, idxs in inv.items():
        for i in idxs:
            pos[i] = palabra
    return " ".join(pos[i] for i in sorted(pos))


def buscar_openalex(q, n, y0, y1, cfg):
    par = {"search": q, "per_page": min(n, 100), "filter": f"publication_year:{y0}-{y1},type:article" + {"en": ",language:en", "es": ",language:es"}.get(cfg.get("idioma"), ""),
           "select": "id,doi,title,publication_year,authorships,primary_location,abstract_inverted_index"}
    if cfg.get("key_openalex"):
        par["api_key"] = cfg["key_openalex"]
    if cfg.get("correo"):
        par["mailto"] = cfg["correo"]
    try:
        d = get_json("https://api.openalex.org/works?" + urllib.parse.urlencode(par))
    except RuntimeError as e:
        if any(c in str(e) for c in ("401", "403", "409", "429")):
            raise RuntimeError("omitido: OpenAlex exige clave gratuita (openalex.org/settings/api)")
        raise
    out = []
    for w in d.get("results", []):
        loc = w.get("primary_location") or {}
        out.append(registro(titulo=w.get("title") or "",
                            autores="; ".join(a["author"]["display_name"] for a in w.get("authorships", [])
                                              if a.get("author", {}).get("display_name")),
                            anio=w.get("publication_year"), revista=(loc.get("source") or {}).get("display_name") or "",
                            doi=w.get("doi") or "", url=loc.get("landing_page_url") or w.get("id", ""),
                            resumen=_abstract_openalex(w.get("abstract_inverted_index")), fuente="OpenAlex"))
    return out


def _primero(v):
    if isinstance(v, list):
        return v[0] if v else ""
    return v or ""


def buscar_base(q, n, y0, y1, cfg):
    par = {"func": "PerformSearch", "query": q, "format": "json", "hits": n}
    st, txt = http_get("https://api.base-search.net/cgi-bin/BaseHttpSearchInterface.fcgi?" + urllib.parse.urlencode(par))
    if st in (401, 403) or not txt.strip().startswith("{"):
        raise RuntimeError("omitido: BASE rechazó la consulta (su IP debe estar autorizada por BASE)")
    docs = json.loads(txt).get("response", {}).get("docs", [])
    out = []
    for r in docs:
        cr = r.get("dccreator", [])
        doi = _primero(r.get("dcdoi")) or (lambda m: m.group(0) if m else "")(DOI_RE.search(json.dumps(r)))
        out.append(registro(titulo=_primero(r.get("dctitle")), autores="; ".join(cr) if isinstance(cr, list) else str(cr),
                            anio=anio_de(r.get("dcyear")), revista=_primero(r.get("dcsource")), doi=doi,
                            url=_primero(r.get("dclink")) or _primero(r.get("dcidentifier")),
                            resumen=_primero(r.get("dcdescription")), fuente="BASE"))
    return out


def buscar_dimensions(q, n, y0, y1, cfg):
    if not cfg.get("key_dimensions"):
        raise RuntimeError("omitido: sin credenciales de Dimensions (API de suscripción/institucional)")
    st, txt = http_post("https://app.dimensions.ai/api/auth.json", json.dumps({"key": cfg["key_dimensions"]}),
                        {"Content-Type": "application/json"})
    if st != 200:
        raise RuntimeError("omitido: Dimensions rechazó las credenciales")
    tok = json.loads(txt)["token"]
    qs = q.replace("\\", " ").replace('"', " ")
    dsl = (f'search publications in title_abstract_only for "{qs}" where year>={y0} and year<={y1} and type="article" '
           f"return publications[doi+title+year+journal+authors+abstract+id] limit {n}")
    st, txt = http_post("https://app.dimensions.ai/api/dsl/v2", dsl, {"Authorization": "JWT " + tok})
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    out = []
    for r in json.loads(txt).get("publications", []):
        au = "; ".join(f"{a.get('last_name', '')}, {a.get('first_name', '')}".strip(", ") for a in r.get("authors", []))
        out.append(registro(titulo=r.get("title", ""), autores=au, anio=r.get("year"),
                            revista=(r.get("journal") or {}).get("title", ""), doi=r.get("doi", ""),
                            url=f"https://app.dimensions.ai/details/publication/{r.get('id', '')}",
                            resumen=r.get("abstract", ""), fuente="Dimensions"))
    return out


def _lista_de_dicts(obj):
    """Localiza, dentro de un JSON arbitrario, la lista de registros (dicts con 'title')."""
    if isinstance(obj, list) and obj and isinstance(obj[0], dict) and any("title" in k.lower() for k in obj[0]):
        return obj
    if isinstance(obj, dict):
        for v in obj.values():
            r = _lista_de_dicts(v)
            if r:
                return r
    return []


def buscar_jisc(q, n, y0, y1, cfg):
    """JISC Library Hub Discover (catálogo colectivo del Reino Unido). A mejor esfuerzo: es un catálogo de
    bibliotecas (libros, tesis) con pocos DOI; si el servicio no devuelve JSON se leen los DOI de la página."""
    url = "https://discover.libraryhub.jisc.ac.uk/search?" + urllib.parse.urlencode({"q": q, "format": "json", "rn": 1})
    st, txt = http_get(url)
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    try:
        filas = _lista_de_dicts(json.loads(txt))[:n]
    except ValueError:
        return _dois_en_html(url, "JISC Library Hub Discover", n, cfg)
    out = []
    for r in filas:
        low = {k.lower(): v for k, v in r.items()}
        t = _primero(low.get("title") or low.get("titles"))
        t = t.get("title", "") if isinstance(t, dict) else t
        au = low.get("author") or low.get("authors") or low.get("creator") or ""
        m = DOI_RE.search(json.dumps(r))
        out.append(registro(titulo=str(t), autores="; ".join(map(str, au)) if isinstance(au, list) else str(au),
                            anio=anio_de(low.get("date") or low.get("year") or low.get("publication_date")),
                            revista=str(_primero(low.get("publisher") or "")), doi=m.group(0) if m else "",
                            url=str(_primero(low.get("url") or low.get("link") or "")), fuente="JISC Library Hub Discover"))
    return out


def buscar_elsevier(q, n, y0, y1, cfg):
    """ScienceDirect Search API v2 (PUT). Sin clave: Crossref con prefijo 10.1016 (Elsevier)."""
    if not cfg.get("key_elsevier"):
        return buscar_crossref(q, n, y0, y1, cfg, prefijo="10.1016", fuente="Elsevier (Crossref, prefijo 10.1016)")
    cuerpo = json.dumps({"qs": q, "pub": "", "display": {"offset": 0, "show": min(n, 100)}})
    st, txt = http_post("https://api.elsevier.com/content/search/sciencedirect", cuerpo,
                        {"X-ELS-APIKey": cfg["key_elsevier"], "Accept": "application/json",
                         "Content-Type": "application/json"})
    if st in (401, 403):
        raise RuntimeError("omitido: Elsevier rechazó la clave de API (dev.elsevier.com)")
    if st == 429:
        raise RuntimeError("omitido: Elsevier superó la cuota semanal de su clave")
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    out = []
    for r in json.loads(txt).get("results", []):
        au = r.get("authors") or []
        if au and isinstance(au[0], dict):
            au = [a.get("name") or f"{a.get('surname', '')}, {a.get('given_name', '')}".strip(", ") for a in au]
        if isinstance(au, str):
            au = [au]
        doi = r.get("doi") or ""
        out.append(registro(titulo=r.get("title", ""), autores="; ".join(map(str, au)),
                            anio=anio_de(r.get("publicationDate") or r.get("coverDate")),
                            revista=r.get("sourceTitle") or r.get("publicationName") or "", doi=doi,
                            url=r.get("uri") or r.get("link") or (f"https://doi.org/{doi}" if doi else ""),
                            fuente="Elsevier (ScienceDirect)"))
    return out


def buscar_scilit(q, n, y0, y1, cfg):
    raise RuntimeError("omitido: Scilit no tiene API pública documentada; use el enlace de búsqueda manual")


def buscar_worldcat(q, n, y0, y1, cfg):
    raise RuntimeError("omitido: WorldCat exige credenciales OCLC (institucionales); use el enlace manual")


# ============================================================== idioma y traducción de consultas
IDIOMAS = {"Cualquiera": "", "Inglés": "en", "Español": "es"}
_EN = set("the of and in to is for with on by are from this that an as at be was were which their".split())
_ES = set("de la el en y los las del que por con para una un se su al es como más sus este esta".split())


def detectar_idioma(texto):
    """Heurística por palabras funcionales. Devuelve 'en', 'es' o '' (indeterminado)."""
    ws = re.findall(r"[a-záéíóúñü]+", (texto or "").lower())
    en, es = sum(w in _EN for w in ws), sum(w in _ES for w in ws)
    if max(en, es) < 2:
        return ""
    if en >= es * 1.5:
        return "en"
    if es >= en * 1.5:
        return "es"
    return ""


def traducir(texto, origen, destino, correo=""):
    """Traducción automática con la API pública de MyMemory (gratuita, con límite diario)."""
    par = {"q": texto[:480], "langpair": f"{origen}|{destino}"}
    if correo:
        par["de"] = correo
    d = get_json("https://api.mymemory.translated.net/get?" + urllib.parse.urlencode(par))
    t = (d.get("responseData") or {}).get("translatedText", "")
    if not t or "MYMEMORY WARNING" in t.upper():
        raise RuntimeError("el traductor alcanzó su límite diario; escriba las consultas en inglés manualmente")
    return html.unescape(t).strip()


def traducir_consultas(consultas, correo="", log=print):
    """Devuelve las traducciones al inglés (solo de consultas que no estén ya en inglés)."""
    out = []
    for q in consultas:
        if detectar_idioma(q) == "en":
            continue
        try:
            t = traducir(q, "es", "en", correo)
            if t and norm(t) != norm(q):
                out.append(t)
        except Exception as e:
            log(f"  No se pudo traducir «{q[:40]}»: {str(e)[:70]}")
            break
        time.sleep(0.4)
    return out


# ============================================================== conectores adicionales
def _limpiar_q(q):
    """Quita signos que rompen consultas tipo Lucene/Solr."""
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s\-]", " ", q or "", flags=re.UNICODE)).strip()


def _html_a_texto(s):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s or ""))).strip()


def buscar_arxiv(q, n, y0, y1, cfg):
    """arXiv (preprints): API Atom oficial. Cada artículo tiene el DOI 10.48550/arXiv.<id> (DataCite)."""
    palabras = sorted({w for w in re.findall(r"[A-Za-z0-9\-]{4,}", q) if w.lower() not in STOP}, key=len, reverse=True)
    if not palabras:
        return []
    xml = ""
    for k in (5, 3):                                   # si 5 términos no dan resultados, se relaja a 3
        sq = " AND ".join(f"all:{w}" for w in palabras[:k])
        st, xml = http_get("https://export.arxiv.org/api/query?" + urllib.parse.urlencode(
            {"search_query": sq, "start": 0, "max_results": n, "sortBy": "relevance"}), {"Accept": "application/atom+xml"})
        if st != 200:
            raise RuntimeError(f"HTTP {st}")
        if "<entry>" in xml:
            break
        time.sleep(3)
    A, X = "{http://www.w3.org/2005/Atom}", "{http://arxiv.org/schemas/atom}"
    out = []
    for e in ET.fromstring(xml).findall(A + "entry"):
        url = (e.findtext(A + "id") or "").strip()
        aid = re.sub(r"v\d+$", "", url.split("/abs/")[-1])
        doi = (e.findtext(X + "doi") or "").strip() or (f"10.48550/arXiv.{aid}" if aid else "")
        out.append(registro(
            titulo=" ".join((e.findtext(A + "title") or "").split()),
            autores="; ".join(a.findtext(A + "name", "") for a in e.findall(A + "author")),
            anio=anio_de(e.findtext(A + "published")), revista=(e.findtext(X + "journal_ref") or "arXiv (preprint)").strip(),
            doi=doi, url=url, resumen=" ".join((e.findtext(A + "summary") or "").split()), fuente="arXiv"))
    time.sleep(3)                                      # arXiv pide ≤ 1 consulta cada 3 s
    return out


def buscar_semanticscholar(q, n, y0, y1, cfg):
    par = {"query": _limpiar_q(q), "limit": min(n, 100), "year": f"{y0}-{y1}",
           "fields": "title,abstract,year,authors,venue,externalIds,url"}
    h = {"x-api-key": cfg["key_semantic"]} if cfg.get("key_semantic") else {}
    try:
        d = get_json("https://api.semanticscholar.org/graph/v1/paper/search?" + urllib.parse.urlencode(par), h)
    except RuntimeError as e:
        if "429" in str(e) or "HTTP 403" in str(e):
            raise RuntimeError("omitido: Semantic Scholar limitó las consultas (use una clave gratuita)")
        raise
    out = [registro(titulo=p.get("title") or "", autores="; ".join(a.get("name", "") for a in p.get("authors", [])),
                    anio=p.get("year"), revista=p.get("venue") or "", doi=(p.get("externalIds") or {}).get("DOI", ""),
                    url=p.get("url", ""), resumen=p.get("abstract") or "", fuente="Semantic Scholar")
           for p in d.get("data", [])]
    time.sleep(1.1)
    return out


def buscar_core(q, n, y0, y1, cfg):
    if not cfg.get("key_core"):
        raise RuntimeError("omitido: CORE exige una clave gratuita (core.ac.uk/services/api)")
    h = {"Authorization": "Bearer " + cfg["key_core"]}
    base = "https://api.core.ac.uk/v3/search/works?"
    try:
        d = get_json(base + urllib.parse.urlencode(
            {"q": f"({_limpiar_q(q)}) AND yearPublished>={y0} AND yearPublished<={y1}", "limit": min(n, 100)}), h)
    except RuntimeError as e:
        if "HTTP 401" in str(e) or "HTTP 403" in str(e):
            raise RuntimeError("omitido: CORE rechazó la clave de API")
        if "HTTP 400" not in str(e):
            raise
        d = get_json(base + urllib.parse.urlencode({"q": _limpiar_q(q), "limit": min(n, 100)}), h)
    out = []
    for r in d.get("results", []):
        rev = (r.get("journals") or [{}])[0].get("title", "") if r.get("journals") else ""
        out.append(registro(titulo=r.get("title") or "", autores="; ".join(a.get("name", "") for a in r.get("authors", [])),
                            anio=r.get("yearPublished"), revista=rev or r.get("publisher") or "CORE", doi=r.get("doi") or "",
                            url=r.get("downloadUrl") or (f"https://core.ac.uk/works/{r['id']}" if r.get("id") else ""),
                            resumen=r.get("abstract") or "", fuente="CORE"))
    return out


def buscar_europepmc(q, n, y0, y1, cfg):
    termino = f"({_limpiar_q(q)}) AND (PUB_YEAR:[{y0} TO {y1}])"
    if cfg.get("idioma") == "en":
        termino += ' AND LANG:"eng"'
    elif cfg.get("idioma") == "es":
        termino += ' AND LANG:"spa"'
    par = {"query": termino, "format": "json", "pageSize": min(n, 100), "resultType": "core"}
    d = get_json("https://www.ebi.ac.uk/europepmc/webservices/rest/search?" + urllib.parse.urlencode(par))
    out = []
    for r in (d.get("resultList") or {}).get("result", []):
        rev = r.get("journalTitle") or ((r.get("journalInfo") or {}).get("journal") or {}).get("title", "")
        url = (f"https://pubmed.ncbi.nlm.nih.gov/{r['pmid']}/" if r.get("pmid")
               else f"https://europepmc.org/article/{r.get('source', 'MED')}/{r.get('id', '')}")
        out.append(registro(titulo=(r.get("title") or "").rstrip("."), autores=r.get("authorString", ""),
                            anio=anio_de(r.get("pubYear")), revista=rev, doi=r.get("doi") or "", url=url,
                            resumen=_html_a_texto(r.get("abstractText")), fuente="Europe PMC"))
    return out


def buscar_opencitations(q, n, y0, y1, cfg):
    """OpenCitations es un ÍNDICE DE CITAS (no busca por tema). Se usan como semilla los DOI que su manuscrito ya cita
    y se recuperan las obras que CITAN a esas referencias (red de citas); los metadatos se completan con Crossref."""
    semillas = sorted(cfg.get("_dois_semilla") or [])[:8]
    if not semillas:
        raise RuntimeError("omitido: OpenCitations no busca por tema; el manuscrito no cita DOI que sirvan de semilla")
    if cfg.get("_oc_hecho"):
        return []                                      # la expansión por citas se hace una sola vez
    cfg["_oc_hecho"] = True
    candidatos = Counter()
    for s in semillas:
        filas = []
        for base in ("https://api.opencitations.net/index/v2", "https://opencitations.net/index/api/v2"):
            try:
                filas = get_json(f"{base}/citations/doi:{urllib.parse.quote(s, safe='/()')}")
                break
            except Exception:
                continue
        for f in filas if isinstance(filas, list) else []:
            for tok in str(f.get("citing", "")).split():
                if tok.startswith("doi:"):
                    candidatos[limpiar_doi(tok[4:])] += 1
        time.sleep(0.3)
    if not candidatos:
        return []
    out = []
    for doi, _ in candidatos.most_common(max(n * 3, 15)):
        try:
            it = get_json("https://api.crossref.org/works/" + urllib.parse.quote(doi, safe="/()"))["message"]
        except Exception:
            continue
        r = _cr_registro(it, "OpenCitations (obras que citan sus referencias)")
        if r["titulo"]:
            out.append(r)
        time.sleep(0.15)
    return out


def buscar_zenodo(q, n, y0, y1, cfg):
    par = {"q": f"({_limpiar_q(q)}) AND publication_date:[{y0}-01-01 TO {y1}-12-31]", "size": min(n, 100),
           "sort": "bestmatch", "type": "publication"}
    d = get_json("https://zenodo.org/api/records?" + urllib.parse.urlencode(par))
    out = []
    for h in (d.get("hits") or {}).get("hits", []):
        md = h.get("metadata", {})
        out.append(registro(titulo=md.get("title", ""), autores="; ".join(c.get("name", "") for c in md.get("creators", [])),
                            anio=anio_de(md.get("publication_date")),
                            revista=((md.get("journal") or {}).get("title") or "Zenodo"),
                            doi=h.get("doi") or md.get("doi") or "", url=(h.get("links") or {}).get("self_html", ""),
                            resumen=_html_a_texto(md.get("description")), fuente="Zenodo"))
    return out


def buscar_figshare(q, n, y0, y1, cfg):
    cuerpo = json.dumps({"search_for": _limpiar_q(q), "page_size": min(n, 100), "published_since": f"{y0}-01-01"})
    st, txt = http_post("https://api.figshare.com/v2/articles/search", cuerpo, {"Content-Type": "application/json"})
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    out = []
    for i, a in enumerate(json.loads(txt)):
        autores, resumen = "", ""
        if i < 10:                                     # el listado no trae autores ni resumen: se piden al detalle
            try:
                det = get_json(f"https://api.figshare.com/v2/articles/{a['id']}")
                autores = "; ".join(x.get("full_name", "") for x in det.get("authors", []))
                resumen = _html_a_texto(det.get("description"))
            except Exception:
                pass
        out.append(registro(titulo=a.get("title", ""), autores=autores, anio=anio_de(a.get("published_date")),
                            revista="Figshare (" + (a.get("defined_type_name") or "repositorio") + ")",
                            doi=a.get("doi") or "", url=a.get("url_public_html", ""), resumen=resumen, fuente="Figshare"))
    return out


def buscar_dryad(q, n, y0, y1, cfg):
    d = get_json("https://datadryad.org/api/v2/search?" + urllib.parse.urlencode({"q": _limpiar_q(q), "per_page": min(n, 100)}))
    out = []
    for r in (d.get("_embedded") or {}).get("stash:datasets", []):
        au = "; ".join(f"{a.get('lastName', '')}, {a.get('firstName', '')}".strip(", ") for a in r.get("authors", []))
        out.append(registro(titulo=r.get("title", ""), autores=au, anio=anio_de(r.get("publicationDate")),
                            revista="Dryad (datos de investigación)", doi=str(r.get("identifier", "")).replace("doi:", ""),
                            url=r.get("sharingLink", ""), resumen=_html_a_texto(r.get("abstract")), fuente="Dryad"))
    return out


def buscar_datacite(q, n, y0, y1, cfg):
    par = {"query": f"({_limpiar_q(q)}) AND publicationYear:[{y0} TO {y1}]", "page[size]": min(n, 100),
           "resource-type-id": "text,journal-article,preprint,dissertation,book,book-chapter,conference-paper,report"}
    d = get_json("https://api.datacite.org/dois?" + urllib.parse.urlencode(par))
    out = []
    for r in d.get("data", []):
        a = r.get("attributes", {})
        desc = next((x.get("description", "") for x in a.get("descriptions", []) if x.get("descriptionType") == "Abstract"),
                    (a.get("descriptions") or [{}])[0].get("description", ""))
        pub = a.get("publisher")
        pub = pub.get("name", "") if isinstance(pub, dict) else (pub or "")
        out.append(registro(titulo=(a.get("titles") or [{}])[0].get("title", ""),
                            autores="; ".join(c.get("name", "") for c in a.get("creators", [])),
                            anio=a.get("publicationYear"), revista=(a.get("container") or {}).get("title") or pub or "DataCite",
                            doi=a.get("doi") or "", url=a.get("url") or "", resumen=_html_a_texto(desc), fuente="DataCite"))
    return out


def buscar_lens(q, n, y0, y1, cfg):
    if not cfg.get("key_lens"):
        raise RuntimeError("omitido: Lens.org exige un token de API (solicítelo en lens.org → Subscriptions)")
    cuerpo = json.dumps({
        "query": {"bool": {"must": [{"query_string": {"query": _limpiar_q(q), "fields": ["title", "abstract", "keyword"]}}],
                           "filter": [{"range": {"year_published": {"gte": y0, "lte": y1}}}]}},
        "size": min(n, 100), "include": ["lens_id", "title", "abstract", "authors", "year_published", "external_ids",
                                         "source", "source_urls"]})
    st, txt = http_post("https://api.lens.org/scholarly/search", cuerpo,
                        {"Authorization": "Bearer " + cfg["key_lens"], "Content-Type": "application/json"})
    if st in (401, 403):
        raise RuntimeError("omitido: Lens.org rechazó el token de API")
    if st == 429:
        raise RuntimeError("omitido: Lens.org superó el límite de consultas")
    if st != 200:
        raise RuntimeError(f"HTTP {st}")
    out = []
    for r in json.loads(txt).get("data", []):
        doi = next((x.get("value", "") for x in r.get("external_ids", []) if x.get("type") == "doi"), "")
        au = "; ".join(f"{a.get('last_name', '')}, {a.get('first_name', '')}".strip(", ") for a in r.get("authors", []))
        urls = r.get("source_urls") or []
        out.append(registro(titulo=r.get("title", ""), autores=au, anio=r.get("year_published"),
                            revista=(r.get("source") or {}).get("title", ""), doi=doi,
                            url=(urls[0].get("url", "") if urls else f"https://www.lens.org/lens/scholar/article/{r.get('lens_id', '')}"),
                            resumen=r.get("abstract") or "", fuente="Lens.org"))
    return out


def buscar_openaire(q, n, y0, y1, cfg):
    """OpenAIRE Graph API v1 (el endpoint antiguo /search/publications es inestable)."""
    base = "https://api.openaire.eu/graph/v1/researchProducts?"
    par = {"search": _limpiar_q(q), "type": "publication", "pageSize": min(n, 100),
           "fromPublicationDate": f"{y0}-01-01", "toPublicationDate": f"{y1}-12-31"}
    try:
        d = get_json(base + urllib.parse.urlencode(par))
    except RuntimeError as e:
        if "HTTP 400" not in str(e):
            raise
        d = get_json(base + urllib.parse.urlencode({k: par[k] for k in ("search", "type", "pageSize")}))
    out = []
    for r in d.get("results", []):
        doi = next((p.get("value", "") for p in (r.get("pids") or []) if str(p.get("scheme", "")).lower() == "doi"), "")
        desc = r.get("descriptions") or []
        urls = next((i.get("urls") for i in (r.get("instances") or []) if i.get("urls")), [])
        out.append(registro(titulo=r.get("mainTitle") or r.get("title") or "",
                            autores="; ".join(a.get("fullName", "") for a in (r.get("authors") or []) if isinstance(a, dict)),
                            anio=anio_de(r.get("publicationDate")), revista=(r.get("container") or {}).get("name", "") or "OpenAIRE",
                            doi=doi, url=(urls[0] if urls else ""),
                            resumen=_html_a_texto(" ".join(desc) if isinstance(desc, list) else str(desc)), fuente="OpenAIRE"))
    return out


def _dspace6(base, fuente, etiqueta, q, n):
    """Repositorios DSpace 6 con API REST pública (DOAB, OAPEN): /rest/search?query=…&expand=metadata."""
    d = get_json(base + "/rest/search?" + urllib.parse.urlencode(
        {"query": _limpiar_q(q), "expand": "metadata", "limit": min(n, 100)}))
    if isinstance(d, dict):
        d = d.get("items") or d.get("results") or []
    out = []
    for it in d:
        md = {}
        for m in it.get("metadata", []):
            md.setdefault(m.get("key", ""), []).append(m.get("value", ""))
        g = lambda *ks: [v for k in ks for v in md.get(k, [])]
        doi = ""
        for k, vs in md.items():
            for v in vs:
                if ("doi" in k.lower() or k.startswith("dc.identifier")) and DOI_RE.search(v or ""):
                    doi = DOI_RE.search(v).group(0)
                    break
            if doi:
                break
        out.append(registro(titulo=(g("dc.title") or [it.get("name", "")])[0],
                            autores="; ".join(g("dc.contributor.author", "dc.contributor.editor")),
                            anio=anio_de((g("dc.date.issued", "dc.date.available") or [""])[0]),
                            revista=(g("publisher.name", "dc.publisher") or [etiqueta])[0], doi=doi,
                            url=(g("dc.identifier.uri") or [""])[0],
                            resumen=" ".join(g("dc.description.abstract", "dc.description")), fuente=fuente))
    return out


def buscar_doab(q, n, y0, y1, cfg):
    return _dspace6("https://directory.doabooks.org", "DOAB", "Libro de acceso abierto (DOAB)", q, n)


def buscar_oapen(q, n, y0, y1, cfg):
    return _dspace6("https://library.oapen.org", "OAPEN", "Libro de acceso abierto (OAPEN)", q, n)


def buscar_plos(q, n, y0, y1, cfg):
    """API de búsqueda de PLOS (Solr). Límite oficial: 10 consultas por minuto."""
    par = {"q": f"everything:({_limpiar_q(q)})", "wt": "json", "rows": min(n, 100),
           "fl": "id,title_display,abstract,author_display,publication_date,journal",
           "fq": f"publication_date:[{y0}-01-01T00:00:00Z TO {y1}-12-31T23:59:59Z]"}
    try:
        d = get_json("https://api.plos.org/search?" + urllib.parse.urlencode(par))
    except RuntimeError:
        return buscar_crossref(q, n, y0, y1, cfg, prefijo="10.1371", fuente="PLOS (Crossref, prefijo 10.1371)")
    out = []
    for r in (d.get("response") or {}).get("docs", []):
        ab = r.get("abstract") or ""
        out.append(registro(titulo=r.get("title_display", ""), autores="; ".join(r.get("author_display") or []),
                            anio=anio_de(r.get("publication_date")), revista=r.get("journal", "PLOS"), doi=r.get("id", ""),
                            url=f"https://doi.org/{r['id']}" if r.get("id") else "",
                            resumen=_html_a_texto(" ".join(ab) if isinstance(ab, list) else ab), fuente="PLOS"))
    time.sleep(6.2)
    return out


def buscar_elife(q, n, y0, y1, cfg):
    """eLife: API oficial de búsqueda; si no responde se usa Crossref con el prefijo DOI 10.7554 (eLife)."""
    try:
        d = get_json("https://api.elifesciences.org/search?" + urllib.parse.urlencode(
            {"for": _limpiar_q(q), "per-page": min(n, 100), "sort": "relevance", "order": "desc", "type[]": "research-article"}),
            {"Accept": "application/vnd.elife.search+json; version=2"})
        out = []
        for it in d.get("items", []):
            doi = it.get("doi") or (f"10.7554/eLife.{it['id']}" if it.get("id") else "")
            out.append(registro(titulo=_html_a_texto(it.get("title")), autores=it.get("authorLine", ""),
                                anio=anio_de(it.get("published") or it.get("statusDate")), revista="eLife", doi=doi,
                                url=f"https://elifesciences.org/articles/{it['id']}" if it.get("id") else "",
                                resumen=_html_a_texto(it.get("impactStatement")), fuente="eLife"))
        if out:
            return out
    except Exception:
        pass
    return buscar_crossref(q, n, y0, y1, cfg, prefijo="10.7554", fuente="eLife (Crossref, prefijo 10.7554)")


FUENTES = {
    "Google Scholar": buscar_scholar,
    "SciELO": buscar_scielo,
    "DOAJ": buscar_doaj,
    "Crossref": lambda q, n, y0, y1, cfg: buscar_crossref(q, n, y0, y1, cfg),
    "SpringerOpen": buscar_springeropen,
    "PubMed": buscar_pubmed,
    "ERIC": buscar_eric,
    "Redalyc": buscar_redalyc,
    "RENATI": buscar_renati,
    "La Referencia": buscar_lareferencia,
    "OpenAlex": buscar_openalex,
    "BASE": buscar_base,
    "Dimensions": buscar_dimensions,
    "JISC": buscar_jisc,
    "KERWA": buscar_kerwa,
    "Elsevier": buscar_elsevier,
    "Scilit": buscar_scilit,
    "WorldCat": buscar_worldcat,
    "arXiv": buscar_arxiv,
    "Semantic Scholar": buscar_semanticscholar,
    "CORE": buscar_core,
    "Europe PMC": buscar_europepmc,
    "OpenCitations": buscar_opencitations,
    "Zenodo": buscar_zenodo,
    "Figshare": buscar_figshare,
    "Dryad": buscar_dryad,
    "DataCite": buscar_datacite,
    "Lens.org": buscar_lens,
    "OpenAIRE": buscar_openaire,
    "DOAB": buscar_doab,
    "OAPEN": buscar_oapen,
    "PLOS": buscar_plos,
    "eLife": buscar_elife,
}

ENLACES_MANUALES = {
    "Google Scholar": "https://scholar.google.com/scholar?q={q}",
    "SciELO": "https://search.scielo.org/?q={q}",
    "DOAJ": "https://doaj.org/search/articles?source=%7B%22query%22%3A%7B%22query_string%22%3A%7B%22query%22%3A%22{q}%22%7D%7D%7D",
    "Crossref": "https://search.crossref.org/?q={q}",
    "SpringerOpen": "https://www.springeropen.com/search?query={q}",
    "PubMed": "https://pubmed.ncbi.nlm.nih.gov/?term={q}",
    "ERIC": "https://eric.ed.gov/?q={q}",
    "Redalyc": "https://www.redalyc.org/busquedaArticuloFiltros.oa?q={q}",
    "RENATI": "https://renati.sunedu.gob.pe/discover?query={q}",
    "La Referencia": "https://www.lareferencia.info/vufind/Search/Results?lookfor={q}&type=AllFields",
    "OpenAlex": "https://openalex.org/works?search={q}",
    "BASE": "https://www.base-search.net/Search/Results?lookfor={q}&type=all",
    "Dimensions": "https://app.dimensions.ai/discover/publication?search_text={q}",
    "JISC": "https://discover.libraryhub.jisc.ac.uk/search?q={q}",
    "KERWA": "https://www.kerwa.ucr.ac.cr/search?query={q}",
    "Elsevier": "https://www.sciencedirect.com/search?qs={q}",
    "Scilit": "https://www.scilit.com/  (pegue allí la consulta: {q})",
    "WorldCat": "https://search.worldcat.org/search?q={q}",
    "arXiv": "https://arxiv.org/search/?query={q}&searchtype=all",
    "Semantic Scholar": "https://www.semanticscholar.org/search?q={q}",
    "CORE": "https://core.ac.uk/search?q={q}",
    "Europe PMC": "https://europepmc.org/search?query={q}",
    "OpenCitations": "https://opencitations.net/  (índice de citas: se consulta por DOI, no por tema)",
    "Zenodo": "https://zenodo.org/search?q={q}",
    "Figshare": "https://figshare.com/search?q={q}",
    "Dryad": "https://datadryad.org/search?q={q}",
    "DataCite": "https://commons.datacite.org/?query={q}",
    "Lens.org": "https://www.lens.org/lens/search/scholar/list?q={q}",
    "OpenAIRE": "https://explore.openaire.eu/search/find?keyword={q}",
    "DOAB": "https://directory.doabooks.org/discover?query={q}",
    "OAPEN": "https://library.oapen.org/discover?query={q}",
    "PLOS": "https://journals.plos.org/plosone/search?q={q}",
    "eLife": "https://elifesciences.org/search?for={q}",
}


# ============================================================== validación (DOI.org + Crossref)
def doi_existe_en_doiorg(doi):
    """DOI.org (Handle System): responseCode 1 = el DOI está registrado."""
    try:
        st, txt = http_get("https://doi.org/api/handles/" + urllib.parse.quote(doi, safe="/()"))
        if st == 404:
            return False
        return json.loads(txt).get("responseCode") == 1
    except Exception:
        return None          # indeterminado (red)


def validar(rec, cfg):
    """Devuelve (ok, motivo)."""
    if rec["doi"]:
        ex = doi_existe_en_doiorg(rec["doi"])
        if ex is False:
            return False, "DOI inexistente en DOI.org"
        if ex is None:
            return False, "no se pudo validar el DOI (error de red)"
        if rec["meta_crossref"]:
            return True, ""
        try:
            it = get_json("https://api.crossref.org/works/" + urllib.parse.quote(rec["doi"], safe="/()"))["message"]
        except Exception:
            return True, ""              # existe en DOI.org pero no en Crossref (p. ej. DataCite)
        cr = _cr_registro(it, "")
        sim = difflib.SequenceMatcher(None, norm(cr["titulo"]), norm(rec["titulo"])).ratio()
        inc = len(tokens(cr["titulo"]) & tokens(rec["titulo"])) / max(1, len(tokens(cr["titulo"])))
        if sim < 0.60 and inc < 0.60:
            return False, "el DOI corresponde a otra obra"
        rec["resumen"] = rec["resumen"] or cr["resumen"]
        rec["autores"] = rec["autores"] or cr["autores"]
        rec["revista"] = rec["revista"] or cr["revista"]
        rec["anio"] = rec["anio"] or cr["anio"]
        return True, ""
    if cfg.get("aceptar_handle") and rec["url"]:
        try:
            st, _ = http_get(rec["url"], timeout=20, reintentos=1)
            return (True, "") if st < 400 else (False, "URL/identificador persistente no accesible")
        except Exception:
            return False, "URL/identificador persistente no accesible"
    return False, "sin DOI verificable"


# ============================================================== pipeline de búsqueda y cribado
def ejecutar(perfil, consultas, fuentes, cfg, y0, y1, n, umbral, tope, log, paso, evento=None):
    ev = evento or (lambda *a: None)      # evento(tipo, ...) alimenta el panel de resultados en proceso
    prisma = dict(por_fuente={}, estados={}, duplicados=0, no_elegibles=0, ya_citados=0, cribados=0,
                  excl_relevancia=0, excl_tope=0, buscados=0, no_recuperados=0, evaluados=0, excl_motivos=Counter(),
                  incluidos=0, consultas=consultas, y0=y0, y1=y1, fuentes=fuentes, umbral=umbral, tope=tope, n_consulta=n,
                  idioma=cfg.get("idioma", ""))
    cfg["_dois_semilla"], cfg["_oc_hecho"] = set(perfil["dois_citados"]), False
    todos = []
    for f in fuentes:
        prisma["por_fuente"][f] = 0
        prisma["estados"][f] = "consultada"
        errores = 0
        ev("fuente", f, "Consultando…", 0)
        for q in consultas:
            try:
                recs = FUENTES[f](q, n, y0, y1, cfg)
                for r in recs:
                    r["consulta"] = q
                todos += recs
                prisma["por_fuente"][f] += len(recs)
                log(f"  {f}: {len(recs):>3} registros ← «{q[:55]}»")
                ev("registros", f, recs)
                ev("fuente", f, "Consultando…", prisma["por_fuente"][f])
            except Exception as e:
                errores += 1
                msg = str(e)
                log(f"  {f}: no disponible ({msg[:80]})")
                if "omitido" in msg or errores >= 2 and prisma["por_fuente"][f] == 0:
                    prisma["estados"][f] = msg[:90]
                    paso(len(consultas) - consultas.index(q))
                    break
            paso(1)
            time.sleep(0.25)
        if errores and prisma["por_fuente"][f] == 0 and prisma["estados"][f] == "consultada":
            prisma["estados"][f] = "sin resultados o no accesible"
        ev("fuente", f, "Consultada" if prisma["estados"][f] == "consultada" else prisma["estados"][f],
           prisma["por_fuente"][f])

    prisma["identificados"] = len(todos)

    # 1) duplicados
    vistos, unicos = {}, []
    for r in todos:
        clave = r["doi"] or norm(r["titulo"])
        if not clave:
            unicos.append(r)
            continue
        if clave in vistos:
            vistos[clave]["fuentes_extra"].add(r["fuente"])
            continue
        r["fuentes_extra"] = {r["fuente"]}
        vistos[clave] = r
        unicos.append(r)
    for r in unicos:
        r.setdefault("fuentes_extra", {r["fuente"]})
    prisma["duplicados"] = len(todos) - len(unicos)
    for r in unicos:
        r.setdefault("decision", "")
        r.setdefault("relevancia", 0.0)

    # 2) no elegibles (sin título o fuera del rango de años)
    elegibles = []
    for r in unicos:
        if not r["titulo"].strip() or (r["anio"] and not (y0 <= r["anio"] <= y1)):
            prisma["no_elegibles"] += 1
            r["decision"] = "excluido: no elegible (sin título o fuera del rango de años)"
        elif cfg.get("idioma") and detectar_idioma(r["titulo"] + " " + r["resumen"]) not in ("", cfg["idioma"]):
            prisma["no_elegibles"] += 1
            r["decision"] = "excluido: no elegible (idioma distinto al seleccionado)"
        else:
            elegibles.append(r)

    # 3) otras razones: ya citadas en el manuscrito
    pendientes = []
    for r in elegibles:
        if r["doi"] and r["doi"] in perfil["dois_citados"]:
            prisma["ya_citados"] += 1
            r["decision"] = "excluido: ya citada en el manuscrito"
        else:
            pendientes.append(r)

    # 4) cribado automático por relevancia
    terminos = set()
    for q in consultas:
        terminos |= tokens(q)
    terminos |= {t for k in perfil["keywords"] for t in tokens(k)}
    denom = max(1, min(len(terminos), 10))
    prisma["cribados"] = len(pendientes)
    for r in pendientes:
        r["relevancia"] = round(min(1.0, len(terminos & tokens(r["titulo"] + " " + r["resumen"])) / denom), 2)
    pendientes.sort(key=lambda r: r["relevancia"], reverse=True)
    pasan = [r for r in pendientes if r["relevancia"] >= umbral]
    for r in pendientes:
        if r["relevancia"] < umbral:
            r["decision"] = "excluido: relevancia insuficiente"
    for r in pasan[tope:]:
        r["decision"] = "excluido: fuera del tope de los más relevantes"
    prisma["excl_relevancia"] = len(pendientes) - len(pasan)
    prisma["excl_tope"] = max(0, len(pasan) - tope)
    pasan = pasan[:tope]
    prisma["buscados"] = len(pasan)

    # 5) recuperación y elegibilidad
    recuperables = []
    for r in pasan:
        if not r["doi"] and not r["url"]:
            prisma["no_recuperados"] += 1
            r["decision"] = "excluido: sin DOI ni URL (no recuperable)"
        else:
            recuperables.append(r)
    prisma["evaluados"] = len(recuperables)
    log(f"Validando {len(recuperables)} informes en DOI.org y Crossref…")
    paso(0, total_extra=len(recuperables))
    incluidos = []
    for r in recuperables:
        ok, motivo = validar(r, cfg)
        r["decision"] = "incluido" if ok else "excluido: " + motivo
        if ok:
            incluidos.append(r)
        else:
            prisma["excl_motivos"][motivo] += 1
        paso(1)
        time.sleep(0.1)
    prisma["incluidos"] = len(incluidos)
    ev("decisiones", todos, incluidos)
    return prisma, incluidos, unicos


# ============================================================== agrupación por hallazgos
TIPOS = [   # (tipo de hallazgo, patrón); se aplica el primero que coincida
    ("Metaanálisis / revisión sistemática", r"meta-?analy|systematic review|umbrella review|scoping review|evidence synthesis|"
                                            r"meta-?an[aá]lisis|revisi[oó]n sistem[aá]tica|revisi[oó]n paraguas"),
    ("Ensayo clínico / estudio experimental", r"randomi[sz]ed|clinical trial|\btrial\b|experimental|ensayo|cuasi-?experiment"),
    ("Validación de instrumento / medida", r"validation|psychometric|\bscale\b|questionnaire|inventory|validaci[oó]n|"
                                            r"\bescala\b|cuestionario|instrumento"),
    ("Tesis / trabajo académico", r"\btesis\b|thesis|dissertation|trabajo de grado|trabajo final"),
    ("Marco teórico / revisión narrativa", r"\breview\b|framework|\btheory\b|theoretical|conceptual|model of|revisi[oó]n|"
                                           r"\bteor[ií]a\b|marco te[oó]rico|perspective"),
    ("Estudio empírico (observacional / cualitativo)", r"cohort|cross-sectional|longitudinal|qualitative|survey|"
                                                        r"correlational|empirical|estudio|study"),
]


def tipo_hallazgo(r):
    txt = f"{r['titulo']} {r['resumen'][:600]} {r['revista']}".lower()
    for nombre, patron in TIPOS:
        if re.search(patron, txt):
            return nombre
    return "Otro / no clasificado"


def temas_del_manuscrito(perfil, consultas):
    """Temas = palabras clave del manuscrito (sin duplicados); si no hay, términos frecuentes."""
    base = list(perfil["keywords"]) or list(perfil["terminos"][:8])
    vistos, temas = set(), []
    for k in base:
        n = norm(k)
        if n and n not in vistos:
            vistos.add(n)
            temas.append(k)
    return temas[:10]


def agrupar_hallazgos(incluidos, perfil, consultas):
    """Asigna a cada referencia el grupo temático (palabra clave del manuscrito) que mejor la describe y su tipo de
    evidencia. Es una clasificación léxica automática: revísela antes de usarla."""
    temas = temas_del_manuscrito(perfil, consultas)
    for r in incluidos:
        mejor, puntaje = "Otros hallazgos relacionados", 0.0
        tt, ta = tokens(r["titulo"], 3), tokens(r["resumen"], 3)
        for t in temas:
            tk = tokens(t, 3)
            if not tk:
                continue
            sc = (2 * len(tk & tt) + len(tk & ta)) / len(tk)
            if sc > puntaje:
                mejor, puntaje = t, sc
        r["grupo"] = mejor if puntaje >= 1.0 else "Otros hallazgos relacionados"
        r["tipo"] = tipo_hallazgo(r)
    grupos = {}
    for r in incluidos:
        grupos.setdefault(r["grupo"], []).append(r)
    orden = sorted(grupos.items(), key=lambda kv: (kv[0] == "Otros hallazgos relacionados", -len(kv[1])))
    return orden


def cita_corta(r):
    au = (r["autores"] or "Autor no disponible").split(";")[0].split(",")[0].strip()
    mas = " et al." if ";" in (r["autores"] or "") else ""
    return f"{au}{mas} ({r['anio'] or 's. f.'})"


# ============================================================== Word único y editable
_ORDEN_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]

MOT_EN = {
    "DOI inexistente en DOI.org": "DOI not found in DOI.org",
    "no se pudo validar el DOI (error de red)": "DOI could not be validated (network error)",
    "el DOI corresponde a otra obra": "DOI belongs to a different work",
    "URL/identificador persistente no accesible": "URL / persistent identifier not accessible",
    "sin DOI verificable": "No verifiable DOI"}


def sombrear(cel, color):
    pr = cel._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    pr.append(s)


def _ordenar_tabla(tabla):
    """Reordena los hijos de cada w:tcPr según el esquema de Word (evita archivos 'dañados')."""
    for tc in tabla._tbl.iter(qn("w:tc")):
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            continue
        hijos = list(tcPr)
        hijos.sort(key=lambda e: _ORDEN_TCPR.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in _ORDEN_TCPR else 99)
        for h in hijos:
            tcPr.remove(h)
        for h in hijos:
            tcPr.append(h)


def _bordes(cel, sz=8, color="000000", quitar=False):
    tcPr = cel._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil" if quitar else "single")
        if not quitar:
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), color)
        b.append(e)
    tcPr.append(b)


def _hipervinculo(par, url, texto, tam=7.5):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(tam * 2))); rpr.append(sz)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = texto; r.append(t)
    h.append(r)
    par._p.append(h)


def guardar_documento(doc, ruta):
    """Guarda el Word; si está abierto/bloqueado prueba nombres alternativos y otras carpetas."""
    carpeta, nombre = os.path.split(os.path.abspath(ruta))
    base, ext = os.path.splitext(nombre)
    destinos = [os.path.join(carpeta, nombre)] + [os.path.join(carpeta, f"{base}_{i}{ext}") for i in range(1, 6)]
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    for otra in (os.path.join(os.path.expanduser("~"), "Documents"), os.path.expanduser("~"), tempfile.gettempdir()):
        destinos.append(os.path.join(otra, f"{base}_{marca}{ext}"))
    ultimo = None
    for d in destinos:
        try:
            doc.save(d)
            return d
        except (PermissionError, OSError) as e:
            ultimo = e
    raise ultimo


# ---- tablas genéricas
def _tabla(doc, cabecera, anchos, filas, tam=8, centrar=(), col_url=None):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in enumerate(anchos):
        t.columns[c].width = Cm(w)
    for c, h in enumerate(cabecera):
        cel = t.rows[0].cells[c]
        cel.width = Cm(anchos[c])
        cel.text = ""
        r = cel.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(tam)
        sombrear(cel, "D9D9D9")
    tr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); tr.append(th)
    for f in filas:
        row = t.add_row()
        for c, v in enumerate(f):
            cel = row.cells[c]
            cel.width = Cm(anchos[c])
            cel.text = ""
            par = cel.paragraphs[0]
            if c in centrar:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            v = "" if v is None else str(v)
            if col_url and c in col_url and v.startswith("http") and " " not in v:
                _hipervinculo(par, v, v, tam)
            else:
                par.add_run(v).font.size = Pt(tam)
    _ordenar_tabla(t)
    return t


# ---- diagrama de flujo editable (estructura del modelo)
def tx(texto, negrita=False, tam=9.0):
    return ([(texto, negrita)], tam)


def ln(etiqueta, valor, tam=9.0):
    return ([(etiqueta, True), (valor, False)], tam)


def _escribir(cel, lineas, centrar=True):
    cel.text = ""
    for i, (runs, tam) in enumerate(lineas):
        par = cel.paragraphs[0] if i == 0 else cel.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrar else WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        for texto, neg in runs:
            r = par.add_run(texto)
            r.bold = neg
            r.font.size = Pt(tam)
    cel.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _bloque(doc, etiqueta_lateral, filas):
    """Bloque del modelo: etiqueta vertical a la izquierda, cuadro de etiqueta y cuadro de contenido.
    filas: ('fila', [lineas_etiqueta], [lineas_contenido]) o ('flecha', texto)."""
    W_LAT, W_ETQ, W_CONT = 1.4, 5.0, 15.8
    expand = []
    for i, f in enumerate(filas):
        if i > 0 and f[0] == "fila" and filas[i - 1][0] == "fila":
            expand.append(("espacio",))
        expand.append(f)

    outer = doc.add_table(rows=1, cols=1)
    outer.style = "Table Grid"
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))     # el bloque no se corta entre páginas
    oc = outer.cell(0, 0)
    oc.width = Cm(W_LAT + W_ETQ + W_CONT + 0.4)
    nt = oc.add_table(rows=len(expand), cols=3)
    nt.alignment = WD_TABLE_ALIGNMENT.CENTER
    nt.autofit = False
    for c, w in enumerate((W_LAT, W_ETQ, W_CONT)):
        nt.columns[c].width = Cm(w)

    for r, f in enumerate(expand):
        row = nt.rows[r]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, w in enumerate((W_LAT, W_ETQ, W_CONT)):
            row.cells[c].width = Cm(w)
        if f[0] == "espacio":
            row.height = Cm(0.15)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            for c in (1, 2):
                row.cells[c].text = ""
                _bordes(row.cells[c], quitar=True)
            continue
        if f[0] == "flecha":
            _bordes(row.cells[1], quitar=True)
            _bordes(row.cells[2], quitar=True)
            row.cells[1].text = ""
            _escribir(row.cells[2], [([("▼  ", False), (f[1], True)], 9.5)])
            continue
        _, etq, cont = f
        for c, lineas in ((1, etq), (2, cont)):
            _bordes(row.cells[c])
            _escribir(row.cells[c], lineas)

    lat = nt.cell(0, 0).merge(nt.cell(len(expand) - 1, 0)) if len(expand) > 1 else nt.cell(0, 0)
    for extra in lat.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    _escribir(lat, [([(etiqueta_lateral, False)], 10.5)])
    for run in lat.paragraphs[0].runs:
        run.italic = True
    tcPr = lat._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), "btLr")                              # texto vertical, de abajo hacia arriba
    tcPr.append(td)
    _bordes(lat, quitar=True)
    _ordenar_tabla(nt)

    primero = oc.paragraphs[0]
    if not primero.text.strip():
        primero._element.getparent().remove(primero._element)
    ultimo = oc.paragraphs[-1]
    ultimo.paragraph_format.space_after = Pt(0)
    ultimo.paragraph_format.line_spacing = Pt(2)
    return outer


def _flecha_grande(doc, texto):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)
    r = par.add_run("⬇   ")
    r.font.size = Pt(20)
    r.font.name = "Segoe UI Symbol"
    r.font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)
    r2 = par.add_run(texto)
    r2.bold = True
    r2.font.size = Pt(10.5)


def _diagrama_editable(doc, p, idioma):
    """Diagrama de flujo PRISMA 2020 con la estructura del modelo (Identificación / Cribado), editable, 'es' o 'en'."""
    en = idioma == "en"
    L = (lambda es, ing: ing if en else es)

    def mot(m):
        return MOT_EN.get(m, m) if en else (m[:1].upper() + m[1:])

    def viñetas(items):
        return [tx(f"• {t}: n = {v}", False, 8.5) for t, v in items]

    # ---------- bloque 1: Identificación
    items = []
    for f in p["fuentes"]:
        marca = "" if p["estados"].get(f) == "consultada" else L(" (no accesible)", " (not accessible)")
        items.append(f"{f}  n = {p['por_fuente'].get(f, 0)}{marca}")
    dato = [tx("   |   ".join(items[i:i + 3]), False, 9) for i in range(0, len(items), 3)]     # 3 fuentes por línea
    ident = [
        ("fila", [tx(L("Base de datos", "Database"), True, 10)],
         [tx(", ".join(p["fuentes"]) + ", DOI.org " + L("(validación de DOI)", "(DOI validation)"), True, 9.5)]),
        ("fila", [tx(L("Método y cadena de búsqueda", "Search method and string"), True, 10)],
         [tx(L("Campos de búsqueda: título, resumen, palabras clave (según cada fuente)",
               "Search fields: title, abstract, keywords (as supported by each source)"), True, 9.5)]
         + [tx(f"{L('Consulta', 'Query')} {i}: {q}", False, 8.5) for i, q in enumerate(p["consultas"], 1)]),
        ("fila", [tx(L("Filtros de búsqueda", "Search filters"), True, 10)],
         [ln(L("Año de publicación: ", "Publication year: "), f"{p['y0']}–{p['y1']}"),
          ln(L("Idioma del documento: ", "Document language: "),
             {"": L("sin restricción (consultas en español o inglés)", "not restricted (queries in Spanish or English)"),
              "en": L("inglés", "English"), "es": L("español", "Spanish")}.get(p.get("idioma", ""), "")),
          ln(L("Tipos de documento: ", "Document types: "),
             L("artículos, actas de congresos, tesis y otros registros académicos con DOI o enlace persistente",
               "articles, conference papers, theses and other scholarly records with DOI or persistent link")),
          ln(L("Relevancia mínima: ", "Minimum relevance: "), f"{p.get('umbral', '—')}"
             + L(f"   |   Tope de referencias: {p.get('tope', '—')}", f"   |   Reference cap: {p.get('tope', '—')}"))]),
        ("fila", [tx(L("Recolección de datos", "Data collection"), True, 10)], dato)]
    _bloque(doc, L("Identificación", "Identification"), ident)

    _flecha_grande(doc, L("Resultados de la búsqueda bibliográfica: n = ", "Literature search results: n = ")
                   + str(p["identificados"]))

    # ---------- bloque 2: Cribado
    quitados = [tx(L("Duplicados eliminados: ", "Duplicates removed: ") + f"n = {p['duplicados']}", False, 9.5),
                tx(L("Marcados como no elegibles (sin título, fuera del rango de años o de otro idioma): ",
                     "Marked as ineligible (no title, outside the year range or other language): ") + f"n = {p['no_elegibles']}",
                   False, 9.5),
                tx(L("Eliminados por otras razones (ya citados en el manuscrito): ",
                     "Removed for other reasons (already cited in the manuscript): ") + f"n = {p['ya_citados']}",
                   False, 9.5)]
    excl_cri = p["excl_relevancia"] + p["excl_tope"]
    elegib = [tx(L("Informes no recuperados (sin DOI ni URL): ", "Reports not retrieved (no DOI or URL): ")
                 + f"n = {p['no_recuperados']}", False, 9.5),
              tx(L("Informes evaluados para elegibilidad (DOI.org + Crossref): ",
                   "Reports assessed for eligibility (DOI.org + Crossref): ") + f"n = {p['evaluados']}", False, 9.5),
              tx(L("Informes excluidos: ", "Reports excluded: ") + f"n = {sum(p['excl_motivos'].values())}", False, 9.5)]
    elegib += viñetas([(mot(m), c) for m, c in p["excl_motivos"].most_common()])
    elegib.append(tx(L("Estudios incluidos tras la evaluación de elegibilidad: ",
                       "Studies included after eligibility assessment: ") + f"n = {p['incluidos']}", True, 10))
    cribado = [
        ("fila", [tx(L("Eliminación de duplicados", "Duplicate removal"), True, 10)], quitados),
        ("flecha", L("Registros para el primer cribado: n = ", "Records for the first screening: n = ")
         + str(p["cribados"])),
        ("fila", [tx(L("Cribado por título y resumen", "Title and abstract screening"), True, 10)],
         [tx(L("Registros excluidos según criterios de exclusión: ", "Records excluded following exclusion criteria: ")
             + f"n = {excl_cri}", False, 9.5)]
         + viñetas([(L("Relevancia insuficiente", "Insufficient relevance"), p["excl_relevancia"]),
                    (L("Fuera del tope de los más relevantes", "Beyond the cap of most relevant records"),
                     p["excl_tope"])])),
        ("flecha", L("Informes para el segundo cribado (recuperación y elegibilidad): n = ",
                     "Reports for the second screening (retrieval and eligibility): n = ") + str(p["buscados"])),
        ("fila", [tx(L("Evaluación de elegibilidad", "Eligibility assessment"), True, 10)], elegib)]
    _bloque(doc, L("Cribado", "Screening"), cribado)

    nota = doc.add_paragraph()
    nota.paragraph_format.space_before = Pt(4)
    r = nota.add_run(L(
        "Diagrama editable: haga clic en cualquier cuadro para modificar el texto o las cifras. Cribado automático "
        "(relevancia léxica y validación de DOI) sin doble revisión humana. Adaptado de Page et al. (2021), BMJ, 372, n71.",
        "Editable diagram: click any box to edit the text or figures. Automatic screening (lexical relevance and DOI "
        "validation) without double human review. Adapted from Page et al. (2021), BMJ, 372, n71."))
    r.italic = True
    r.font.size = Pt(8)


# ---- documento único
def generar_word_unico(ruta, archivo, perfil, p, incluidos, todos):
    """Un solo Word editable con TODOS los resultados: consultas, repositorios, diagramas PRISMA 2020 (ES/EN),
    cuadro por hallazgos, referencias detalladas, cribado completo y limitaciones."""
    d = Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    for k in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, k, Cm(1.8))
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)

    d.add_heading("Resultados de la búsqueda bibliográfica y diagrama PRISMA 2020", 1)
    para = d.add_paragraph()
    for etq, val in (("Manuscrito: ", archivo), ("Título detectado: ", perfil["titulo"] or "—"),
                     ("Fecha: ", f"{datetime.now():%d/%m/%Y %H:%M}"), ("Rango de años: ", f"{p['y0']}–{p['y1']}"),
                     ("Idioma de los documentos: ", {"": "cualquiera", "en": "inglés", "es": "español"}.get(p.get("idioma", ""), "")),
                     ("Desarrolladores del software: ", DESARROLLADORES),
                     ("Registros identificados: ", str(p["identificados"])),
                     ("Referencias nuevas verificadas: ", str(p["incluidos"]))):
        para.add_run(etq).bold = True
        para.add_run(val + "\n")

    d.add_heading("1. Consultas utilizadas", 2)
    for q in p["consultas"]:
        d.add_paragraph(q, style="List Bullet")

    d.add_heading("2. Registros por repositorio", 2)
    q0 = urllib.parse.quote(p["consultas"][0]) if p["consultas"] else ""
    filas = []
    for f in p["fuentes"]:
        inc = sum(1 for r in incluidos if any(x.startswith(f.split(" (")[0]) for x in r["fuentes_extra"]))
        ok = p["estados"].get(f) == "consultada"
        filas.append([f, p["por_fuente"].get(f, 0), inc, "Consultada" if ok else p["estados"].get(f, ""),
                      ENLACES_MANUALES[f].format(q=q0)])
    filas.append(["DOI.org", "—", "—", "Validación de DOI (no es buscador)", "https://doi.org/"])
    _tabla(d, ["Repositorio", "Registros recuperados", "Referencias incluidas", "Estado", "Enlace para búsqueda manual"],
           [3.4, 2.4, 2.4, 6.0, 11.9], filas, 8, centrar=(1, 2), col_url=(4,))

    d.add_page_break()
    d.add_heading("3. Diagrama de flujo PRISMA 2020 (editable) — Español", 2)
    _diagrama_editable(d, p, "es")
    d.add_page_break()
    d.add_heading("3. PRISMA 2020 flow diagram (editable) — English", 2)
    _diagrama_editable(d, p, "en")

    grupos = agrupar_hallazgos(incluidos, perfil, p["consultas"])
    d.add_page_break()
    d.add_heading("4. Cuadro resumen de hallazgos (agrupados por tema del manuscrito)", 2)
    if not grupos:
        d.add_paragraph("No hay referencias incluidas para resumir.")
    else:
        filas = []
        for k, (g, lst) in enumerate(grupos, 1):
            tipos = Counter(r["tipo"] for r in lst)
            anios = [r["anio"] for r in lst if r["anio"]]
            top = sorted(lst, key=lambda r: r["relevancia"], reverse=True)
            clave = "\n".join(f"• {cita_corta(r)}: {r['titulo'][:110]}" for r in top[:3])
            ejemplo = (top[0]["resumen"] or "").strip().replace("\n", " ")
            filas.append([k, g, len(lst), "\n".join(f"{t}: {n}" for t, n in tipos.most_common()),
                          f"{min(anios)}–{max(anios)}" if anios else "—", clave,
                          (ejemplo[:350] + "…") if len(ejemplo) > 350 else (ejemplo or "Sin resumen en la fuente")])
        _tabla(d, ["N.º", "Grupo temático (hallazgo)", "N.º de refs.", "Tipos de evidencia", "Años",
                   "Referencias más relevantes", "Resumen del registro más relevante (según la fuente)"],
               [0.9, 3.4, 1.4, 3.8, 1.6, 8.0, 7.0], filas, 8, centrar=(0, 2, 4))
        n = d.add_paragraph("Nota: la agrupación es léxica y automática (palabras clave del manuscrito); el resumen citado "
                            "es el de la fuente, no una síntesis propia. Lea cada referencia antes de usarla.")
        n.runs[0].italic = True

        d.add_heading(f"5. Hallazgos detallados: referencias nuevas candidatas (n = {len(incluidos)})", 2)
        num = 0
        for g, lst in grupos:
            h = d.add_heading(f"HALLAZGO: {g}   (n = {len(lst)})", 3)
            filas = []
            for r in sorted(lst, key=lambda r: r["relevancia"], reverse=True):
                num += 1
                enlace = f"https://doi.org/{r['doi']}" if r["doi"] else (r["url"] or "—")
                res = (r["resumen"] or "Sin abstract en la fuente").strip()
                filas.append([num, enlace, r["anio"] or "", r["autores"] or "—", r["titulo"], r["revista"] or "—",
                              ", ".join(sorted(r["fuentes_extra"]))[:40], f"{r['relevancia']:.2f}", r["tipo"],
                              (res[:900] + "…") if len(res) > 900 else res])
            _tabla(d, ["N.º", "DOI / enlace", "Año", "Autores", "Título", "Revista", "Fuente", "Relev.",
                       "Tipo de evidencia", "Abstract"],
                   [0.8, 3.3, 1.0, 3.2, 4.4, 2.6, 1.8, 1.0, 2.2, 5.8], filas, 7, centrar=(0, 2, 7), col_url=(1,))

    d.add_page_break()
    d.add_heading("6. Cribado completo (todas las decisiones)", 2)
    filas = [[r["titulo"], r["autores"] or "—", r["anio"] or "", r["doi"] or "—",
              ", ".join(sorted(r["fuentes_extra"])), r.get("relevancia", ""), r["decision"] or "—"] for r in todos]
    _tabla(d, ["Título", "Autores", "Año", "DOI", "Fuente(s)", "Relev.", "Decisión"],
           [7.4, 4.4, 1.1, 3.6, 3.4, 1.2, 5.0], filas, 7, centrar=(2, 5))

    d.add_heading("7. Limitaciones", 2)
    for x in ["arXiv, Zenodo, Figshare, Dryad y DataCite incluyen preprints, informes y conjuntos de datos que NO han pasado "
              "revisión por pares; OpenCitations solo amplía la búsqueda a las obras que citan sus referencias. Verifique el "
              "tipo de documento antes de citar.",
              "El cribado es automático (relevancia léxica + validación de DOI); no hay doble revisión humana ni evaluación "
              "de calidad metodológica. Lea cada referencia antes de citarla.",
              "La relevancia compara palabras de sus consultas con título y resumen; puede perder trabajos en otro idioma.",
              "SciELO, Redalyc, RENATI y La Referencia se consultan a mejor esfuerzo; un resultado vacío no prueba que no "
              "existan trabajos. Para las fuentes sin API use los enlaces de la sección 2.",
              "Para citar el diagrama como parte de una revisión sistemática, registre el protocolo y haga el cribado por pares "
              "(Page et al., 2021, https://doi.org/10.1136/bmj.n71)."]:
        d.add_paragraph(x, style="List Bullet")
    pie = d.add_paragraph()
    pie.paragraph_format.space_before = Pt(8)
    rr = pie.add_run(f"Software desarrollado por {DESARROLLADORES}.")
    rr.italic = True
    rr.font.size = Pt(8.5)
    return guardar_documento(d, ruta)


# ============================================================== interfaz gráfica
def gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, scrolledtext

    cola = queue.Queue()
    est = {"ruta": "", "perfil": None, "ocupado": False}
    raiz = tk.Tk()
    raiz.title("Buscador bibliográfico + diagrama PRISMA 2020 — " + DESARROLLADORES)
    sw, sh = raiz.winfo_screenwidth(), raiz.winfo_screenheight()
    raiz.geometry(f"{min(1560, sw - 40)}x{min(960, sh - 90)}+10+10")

    # ---- dos paneles: controles a la IZQUIERDA (con desplazamiento) y resultados en proceso a la DERECHA
    pw = ttk.PanedWindow(raiz, orient="horizontal")
    pw.pack(fill="both", expand=True)
    izq = ttk.Frame(pw)
    der = ttk.Frame(pw, padding=(8, 8, 8, 4))
    pw.add(izq, weight=1)
    pw.add(der, weight=1)
    cv = tk.Canvas(izq, highlightthickness=0, width=800)
    sb = ttk.Scrollbar(izq, orient="vertical", command=cv.yview)
    cv.configure(yscrollcommand=sb.set)
    sb.pack(side="right", fill="y")
    cv.pack(side="left", fill="both", expand=True)
    m = ttk.Frame(cv, padding=10)
    ventana = cv.create_window((0, 0), window=m, anchor="nw")
    m.bind("<Configure>", lambda e: cv.configure(scrollregion=cv.bbox("all")))
    cv.bind("<Configure>", lambda e: cv.itemconfigure(ventana, width=e.width))

    def rueda(e):
        wd = raiz.winfo_containing(e.x_root, e.y_root)
        while wd is not None and wd is not izq and wd is not raiz:
            if wd.winfo_class() in ("Text", "Treeview", "TCombobox"):
                return
            wd = wd.master
        if wd is izq:
            cv.yview_scroll(int(-e.delta / 120), "units")
    raiz.bind_all("<MouseWheel>", rueda)

    # ---- panel derecho: RESULTADOS EN PROCESO
    ttk.Label(der, text="Resultados en proceso", font=("Segoe UI", 12, "bold")).pack(anchor="w")
    v_cont = tk.StringVar(value="Sin búsqueda en curso. Suba su Word y pulse «Buscar».")
    ttk.Label(der, textvariable=v_cont, foreground="#333333", wraplength=640, justify="left").pack(anchor="w", pady=(0, 6))
    ttk.Label(der, text="Estado por repositorio", font=("Segoe UI", 9, "bold")).pack(anchor="w")
    fr_f = ttk.Frame(der); fr_f.pack(fill="x")
    tv_f = ttk.Treeview(fr_f, columns=("fuente", "estado", "regs", "inc"), show="headings", height=9)
    for c, t, ancho in (("fuente", "Repositorio", 125), ("estado", "Estado", 260), ("regs", "Registros", 85), ("inc", "Incluidos", 80)):
        tv_f.heading(c, text=t)
        tv_f.column(c, width=ancho, anchor="w" if c in ("fuente", "estado") else "center", stretch=(c == "estado"))
    sb_f = ttk.Scrollbar(fr_f, orient="vertical", command=tv_f.yview)
    tv_f.configure(yscrollcommand=sb_f.set)
    tv_f.pack(side="left", fill="x", expand=True); sb_f.pack(side="left", fill="y")
    tv_f.tag_configure("wip", background="#FFF2CC")
    tv_f.tag_configure("ok", background="#E2EFDA")
    tv_f.tag_configure("omit", background="#F2F2F2", foreground="#8A1C1C")
    tv_f.tag_configure("pend", background="#FFFFFF", foreground="#777777")

    ttk.Label(der, text="Registros encontrados (se añaden a medida que llegan)", font=("Segoe UI", 9, "bold")).pack(
        anchor="w", pady=(8, 0))
    fr_r = ttk.Frame(der); fr_r.pack(fill="both", expand=True)
    tv_r = ttk.Treeview(fr_r, columns=("fuente", "anio", "titulo", "decision"), show="headings")
    for c, t, ancho in (("fuente", "Fuente", 105), ("anio", "Año", 45), ("titulo", "Título", 330), ("decision", "Decisión", 190)):
        tv_r.heading(c, text=t)
        tv_r.column(c, width=ancho, anchor="center" if c == "anio" else "w", stretch=(c == "titulo"))
    sb_r = ttk.Scrollbar(fr_r, orient="vertical", command=tv_r.yview)
    sbh_r = ttk.Scrollbar(der, orient="horizontal", command=tv_r.xview)
    tv_r.configure(yscrollcommand=sb_r.set, xscrollcommand=sbh_r.set)
    tv_r.pack(side="left", fill="both", expand=True); sb_r.pack(side="left", fill="y")
    sbh_r.pack(fill="x")
    tv_r.tag_configure("pend", background="#FFFFFF")
    tv_r.tag_configure("inc", background="#E2EFDA")
    tv_r.tag_configure("exc", background="#FADBD8")
    tv_r.tag_configure("dup", background="#EDEDED", foreground="#555555")
    vivo = {"filas": {}, "refs": [], "n": 0, "listas": 0, "total": 0}

    def etiqueta(t):
        ttk.Label(m, text=t, font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(8, 2))

    etiqueta("1. Manuscrito (.docx)")
    f1 = ttk.Frame(m); f1.pack(fill="x")
    v_ruta = tk.StringVar(value="Ningún archivo seleccionado")
    ttk.Label(f1, textvariable=v_ruta, relief="sunken", padding=3).pack(side="left", fill="x", expand=True)
    b_sub = ttk.Button(f1, text="Subir Word y proponer consultas")
    b_sub.pack(side="left", padx=6)

    etiqueta("2. Consultas de búsqueda (una por línea; puede editarlas, también en inglés)")
    t_q = scrolledtext.ScrolledText(m, height=7, font=("Consolas", 10))
    t_q.pack(fill="x")
    fq = ttk.Frame(m); fq.pack(fill="x", pady=(3, 0))
    b_trad = ttk.Button(fq, text="Añadir consultas en inglés (traducción automática)")
    b_trad.pack(side="left")
    ttk.Label(fq, text="  También puede escribirlas usted mismo en inglés.").pack(side="left")

    etiqueta("3. Parámetros")
    f2 = ttk.Frame(m); f2.pack(fill="x")
    v_y0, v_y1 = tk.StringVar(value=str(ANIO_ACTUAL - 10)), tk.StringVar(value=str(ANIO_ACTUAL))
    v_n, v_umb, v_tope = tk.StringVar(value="15"), tk.StringVar(value="0.30"), tk.StringVar(value="120")
    for i, (txt, var, w) in enumerate([("Año desde", v_y0, 6), ("hasta", v_y1, 6), ("Resultados/fuente/consulta", v_n, 5),
                                       ("Relevancia mín. (0–1)", v_umb, 5), ("Tope de referencias", v_tope, 6)]):
        ttk.Label(f2, text=txt).grid(row=i // 3, column=2 * (i % 3), sticky="e", padx=(8, 2), pady=1)
        ttk.Entry(f2, textvariable=var, width=w).grid(row=i // 3, column=2 * (i % 3) + 1, sticky="w")
    ttk.Label(f2, text="Idioma de los documentos").grid(row=2, column=0, columnspan=2, sticky="e", padx=(8, 2), pady=(5, 0))
    v_idioma = tk.StringVar(value="Cualquiera")
    ttk.Combobox(f2, textvariable=v_idioma, values=list(IDIOMAS), state="readonly", width=12).grid(
        row=2, column=2, columnspan=2, sticky="w", pady=(5, 0))

    etiqueta("4. Fuentes")
    f3 = ttk.Frame(m); f3.pack(fill="x")
    v_f = {}
    for i, f in enumerate(FUENTES):
        v_f[f] = tk.BooleanVar(value=True)
        ttk.Checkbutton(f3, text=f, variable=v_f[f]).grid(row=i // 5, column=i % 5, sticky="w", padx=8)

    etiqueta("5. Opcional: correo y claves de API")
    f4 = ttk.Frame(m); f4.pack(fill="x")
    v_mail, v_ks, v_kg = tk.StringVar(), tk.StringVar(), tk.StringVar()
    v_ko, v_kd, v_ke = tk.StringVar(), tk.StringVar(), tk.StringVar()
    v_ksem, v_kcore, v_klens = tk.StringVar(), tk.StringVar(), tk.StringVar()
    for i, (txt, var) in enumerate([("Correo (Crossref/OpenAlex)", v_mail), ("Clave Springer Nature", v_ks),
                                    ("Clave SerpAPI (Google Scholar)", v_kg), ("Clave OpenAlex (gratuita, obligatoria)", v_ko),
                                    ("Credencial Dimensions API", v_kd),
                                    ("Clave Elsevier (dev.elsevier.com)", v_ke),
                                    ("Clave Semantic Scholar (opcional)", v_ksem),
                                    ("Clave CORE (gratuita, obligatoria)", v_kcore),
                                    ("Token Lens.org (obligatorio)", v_klens)]):
        ttk.Label(f4, text=txt).grid(row=i, column=0, sticky="e", padx=4, pady=1)
        ttk.Entry(f4, textvariable=var, width=48).grid(row=i, column=1, sticky="w")
    v_h = tk.BooleanVar(value=False)
    ttk.Checkbutton(m, text="Aceptar registros sin DOI con enlace persistente accesible (p. ej. tesis de RENATI)",
                    variable=v_h).pack(anchor="w", pady=4)

    # ---- botones: buscar / nueva consulta / salir
    fb = ttk.Frame(m); fb.pack(fill="x", pady=6)
    b_go = ttk.Button(fb, text="6. Buscar y generar Word con PRISMA 2020")
    b_go.pack(side="left")
    b_new = ttk.Button(fb, text="Nueva consulta")
    b_new.pack(side="left", padx=(12, 0))
    b_exit = ttk.Button(fb, text="Salir")
    b_exit.pack(side="left", padx=(12, 0))

    barra = ttk.Progressbar(m, mode="determinate"); barra.pack(fill="x")
    v_est = tk.StringVar(value="Listo.")
    ttk.Label(m, textvariable=v_est).pack(anchor="w")
    log = scrolledtext.ScrolledText(m, height=8, state="disabled", font=("Consolas", 9))
    log.pack(fill="both", expand=True, pady=4)
    ttk.Label(m, text=f"Desarrolladores: {DESARROLLADORES}", foreground="#555555").pack(anchor="e")

    def w(msg):
        log.configure(state="normal"); log.insert("end", msg + "\n"); log.see("end"); log.configure(state="disabled")

    LIMITE_FILAS = 6000

    def limpiar_vivo(fuentes=()):
        """Vacía el panel derecho; si se indican fuentes, las lista como «Pendiente»."""
        tv_f.delete(*tv_f.get_children())
        tv_r.delete(*tv_r.get_children())
        vivo.update(filas={}, refs=[], n=0, listas=0, total=len(fuentes))
        for f in fuentes:
            tv_f.insert("", "end", iid=f, values=(f, "Pendiente", "", ""), tags=("pend",))
        v_cont.set("Sin búsqueda en curso. Suba su Word y pulse «Buscar»." if not fuentes else "Preparando la búsqueda…")

    def resumen_vivo(extra=""):
        v_cont.set(f"Registros recuperados: {vivo['n']}   |   Repositorios terminados: {vivo['listas']}/{vivo['total']}" + extra)

    def manejar(a):
        """Procesa un evento del motor de búsqueda y actualiza el panel derecho."""
        tipo = a[0]
        if tipo == "fuente":
            _, f, estado, n_f = a
            if estado not in ("Consultada", "Consultando…") and not estado.startswith(("omitido", "sin ")):
                estado = "No disponible: " + estado
            if tv_f.exists(f):
                final = estado != "Consultando…"
                tag = "wip" if not final else ("ok" if estado == "Consultada" else "omit")
                if final and "listo" not in tv_f.item(f, "tags"):
                    vivo["listas"] += 1
                tv_f.item(f, values=(f, estado, n_f, tv_f.set(f, "inc")), tags=(tag, "listo") if final else (tag,))
                tv_f.see(f)
            resumen_vivo()
        elif tipo == "registros":
            _, f, recs = a
            for r in recs:
                vivo["n"] += 1
                if len(vivo["filas"]) >= LIMITE_FILAS:
                    continue
                iid = tv_r.insert("", "end", values=(f, r["anio"] or "", r["titulo"][:160], "en proceso…"), tags=("pend",))
                vivo["filas"][id(r)] = iid
                vivo["refs"].append(r)
            tv_r.yview_moveto(1.0)
            resumen_vivo()
        elif tipo == "decisiones":
            _, todos, incluidos = a
            ids_inc = {id(r) for r in incluidos}
            n_inc = n_exc = n_dup = 0
            for r in todos:
                iid = vivo["filas"].get(id(r))
                if id(r) in ids_inc:
                    n_inc += 1; dec, tag = "Incluido", "inc"
                elif not r.get("decision"):
                    n_dup += 1; dec, tag = "Duplicado", "dup"
                else:
                    n_exc += 1; dec, tag = r["decision"].replace("excluido: ", "Excluido: ")[:80], "exc"
                if iid and tv_r.exists(iid):
                    tv_r.item(iid, values=(*tv_r.item(iid, "values")[:3], dec), tags=(tag,))
            por_f = Counter(x for r in incluidos for x in r["fuentes_extra"])
            for f in tv_f.get_children():
                tv_f.set(f, "inc", por_f.get(f, 0))
            resumen_vivo(f"   |   Incluidos: {n_inc}   |   Excluidos: {n_exc}   |   Duplicados: {n_dup}")

    def liberar():
        est["ocupado"] = False
        for b in (b_go, b_sub, b_new):
            b.state(["!disabled"])

    def subir():
        r = filedialog.askopenfilename(filetypes=[("Documento Word", "*.docx")])
        if not r:
            return
        try:
            perfil = perfil_manuscrito(leer_parrafos(r))
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo leer el Word: {e}")
            return
        est["ruta"], est["perfil"] = r, perfil
        v_ruta.set(r)
        t_q.delete("1.0", "end")
        t_q.insert("1.0", "\n".join(construir_consultas(perfil)))
        w(f"Título: {perfil['titulo'][:100]}")
        w(f"Palabras clave: {', '.join(perfil['keywords']) or '(no detectadas; se usaron términos frecuentes)'}")
        w(f"DOI ya citados en el manuscrito: {len(perfil['dois_citados'])}")

    def trabajo(salida, consultas, fuentes, cfg, y0, y1, n, umb, tope):
        try:
            total = len(consultas) * len(fuentes)
            cola.put(("max", total))

            def paso(k, total_extra=None):
                if total_extra:
                    cola.put(("addmax", total_extra))
                elif k:
                    cola.put(("paso", k))

            p, inc, todos = ejecutar(est["perfil"], consultas, fuentes, cfg, y0, y1, n, umb, tope,
                                     lambda s: cola.put(("log", s)), paso, lambda *a: cola.put(("ev", a)))
            cola.put(("log", "Generando el Word único (diagramas PRISMA en español e inglés, hallazgos, cribado)…"))
            final = generar_word_unico(salida, os.path.basename(est["ruta"]), est["perfil"], p, inc, todos)
            if os.path.abspath(final) != os.path.abspath(salida):
                cola.put(("log", "Aviso: no se pudo guardar con el nombre elegido (¿está abierto en Word?). "
                                 "Se guardó como: " + final))
            cola.put(("fin", (final, p["incluidos"], p["identificados"])))
        except Exception as e:
            cola.put(("error", str(e)))

    def ir():
        if est["ocupado"]:
            return
        if not est["perfil"]:
            messagebox.showwarning("Falta el manuscrito", "Suba primero el Word.")
            return
        consultas = [q.strip() for q in t_q.get("1.0", "end").splitlines() if q.strip()]
        fuentes = [f for f, v in v_f.items() if v.get()]
        if not consultas or not fuentes:
            messagebox.showwarning("Faltan datos", "Ingrese al menos una consulta y una fuente.")
            return
        if IDIOMAS.get(v_idioma.get()) == "en" and detectar_idioma(" ".join(consultas)) == "es":
            if messagebox.askyesno("Consultas en español", "Eligió documentos en INGLÉS pero las consultas parecen estar en "
                                                           "español.\n\n¿Añadir su traducción al inglés automáticamente?"):
                consultas += [t for t in traducir_consultas(consultas, v_mail.get().strip(), log=w)
                              if norm(t) not in {norm(q) for q in consultas}]
                t_q.delete("1.0", "end")
                t_q.insert("1.0", "\n".join(consultas))
        try:
            y0, y1, n, umb, tope = int(v_y0.get()), int(v_y1.get()), int(v_n.get()), float(v_umb.get()), int(v_tope.get())
        except ValueError:
            messagebox.showwarning("Parámetros", "Revise los valores numéricos.")
            return
        salida = filedialog.asksaveasfilename(title="Guardar el Word con todos los resultados", defaultextension=".docx",
                                              initialfile="Resultados_busqueda_PRISMA_2020.docx",
                                              filetypes=[("Documento Word", "*.docx")])
        if not salida:
            return
        cfg = dict(correo=v_mail.get().strip(), key_springer=v_ks.get().strip(), key_serpapi=v_kg.get().strip(),
                   key_openalex=v_ko.get().strip(), key_dimensions=v_kd.get().strip(),
                   key_elsevier=v_ke.get().strip(), key_semantic=v_ksem.get().strip(), key_core=v_kcore.get().strip(),
                   key_lens=v_klens.get().strip(), idioma=IDIOMAS.get(v_idioma.get(), ""),
                   aceptar_handle=v_h.get())
        est["ocupado"] = True
        for b in (b_go, b_sub, b_new):
            b.state(["disabled"])
        barra["value"] = 0
        limpiar_vivo(fuentes)
        v_est.set("Buscando… esto puede tardar varios minutos.")
        threading.Thread(target=trabajo, args=(salida, consultas, fuentes, cfg, y0, y1, n, umb, tope), daemon=True).start()
        raiz.after(150, revisar)

    def nueva_consulta():
        """Deja la ventana lista para otra búsqueda (conserva correo y claves de API)."""
        if est["ocupado"]:
            messagebox.showinfo("Búsqueda en curso", "Espere a que termine la búsqueda actual.")
            return
        if not messagebox.askyesno("Nueva consulta", "¿Iniciar una nueva consulta?\n\nSe limpiarán el manuscrito, "
                                                     "las consultas y el registro. Los resultados ya guardados no se borran."):
            return
        est["ruta"], est["perfil"] = "", None
        v_ruta.set("Ningún archivo seleccionado")
        t_q.delete("1.0", "end")
        v_y0.set(str(ANIO_ACTUAL - 10)); v_y1.set(str(ANIO_ACTUAL))
        v_n.set("15"); v_umb.set("0.30"); v_tope.set("120")
        for var in v_f.values():
            var.set(True)
        v_h.set(False)
        v_idioma.set("Cualquiera")
        limpiar_vivo()
        barra["value"] = 0
        log.configure(state="normal"); log.delete("1.0", "end"); log.configure(state="disabled")
        v_est.set("Listo para una nueva consulta: suba un Word.")

    def traducir_ui():
        consultas = [q.strip() for q in t_q.get("1.0", "end").splitlines() if q.strip()]
        if not consultas:
            messagebox.showinfo("Sin consultas", "Suba primero el Word o escriba consultas.")
            return
        v_est.set("Traduciendo consultas…")
        raiz.update_idletasks()
        existentes = {norm(q) for q in consultas}
        nuevas = [t for t in traducir_consultas(consultas, v_mail.get().strip(), log=w) if norm(t) not in existentes]
        if nuevas:
            t_q.insert("end-1c", "\n" + "\n".join(nuevas))
            w(f"Se añadieron {len(nuevas)} consultas en inglés.")
        else:
            w("No se añadieron consultas (ya están en inglés o el traductor no respondió).")
        v_est.set("Listo.")

    def salir():
        if est["ocupado"] and not messagebox.askyesno("Salir", "Hay una búsqueda en curso. ¿Salir de todos modos?"):
            return
        raiz.destroy()

    def revisar():
        try:
            while True:
                tipo, x = cola.get_nowait()
                if tipo == "log":
                    w(x)
                elif tipo == "max":
                    barra["maximum"] = x
                elif tipo == "ev":
                    manejar(x)
                elif tipo == "addmax":
                    barra["maximum"] = barra["maximum"] + x
                elif tipo == "paso":
                    barra["value"] = min(barra["maximum"], barra["value"] + x)
                elif tipo == "error":
                    liberar()
                    v_est.set("Error."); messagebox.showerror("Error", x); return
                elif tipo == "fin":
                    liberar()
                    archivo, inc, ident = x
                    v_est.set(f"Terminado: {inc} referencias candidatas de {ident} registros.")
                    w(f"\nWord guardado en: {archivo}")
                    if messagebox.askyesno("Listo", f"Referencias candidatas: {inc}\nRegistros identificados: {ident}\n\n"
                                                    f"{archivo}\n\n¿Abrir el documento ahora?"):
                        try:
                            os.startfile(archivo)          # Windows
                        except Exception:
                            pass
                    return
        except queue.Empty:
            pass
        if est["ocupado"]:
            raiz.after(150, revisar)

    b_sub.configure(command=subir)
    b_trad.configure(command=traducir_ui)
    b_go.configure(command=ir)
    b_new.configure(command=nueva_consulta)
    b_exit.configure(command=salir)
    raiz.protocol("WM_DELETE_WINDOW", salir)
    raiz.mainloop()


if __name__ == "__main__":
    gui()
